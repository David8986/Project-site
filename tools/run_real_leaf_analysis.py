"""Run the real leaf camera/spectrometer analysis.

Inputs live in:
    C:/Users/david/OneDrive/Desktop/data analisys

Outputs are written under:
    outputs/real_leaf_analysis
"""

from __future__ import annotations

import csv
import json
import math
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

import matplotlib.pyplot as plt
import numpy as np
from PIL import Image, ImageDraw

try:
    import cv2
except ModuleNotFoundError:  # pragma: no cover - runtime fallback
    cv2 = None

try:
    from openpyxl import Workbook
    from openpyxl.chart import LineChart, Reference
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ModuleNotFoundError:  # pragma: no cover - runtime fallback
    Workbook = None
    LineChart = Reference = Alignment = Font = PatternFill = get_column_letter = None


PROJECT_ROOT = Path(r"C:\Users\david\OneDrive\Desktop\Archive\Projects\CodeX")
DATA_ROOT = Path(r"C:\Users\david\OneDrive\Desktop\data analisys")
OUT_DIR = PROJECT_ROOT / "outputs" / "real_leaf_analysis"
TARGET_WAVELENGTHS = (532, 556, 680, 725, 850, 940)


@dataclass
class Spectrum:
    path: Path
    relative_path: str
    environment: str
    sample_id: str
    condition: str
    role: str
    modality: str
    filter_label: str
    filter_nm: float | None
    wavelengths: np.ndarray
    intensities: np.ndarray


@dataclass
class CameraMeasurement:
    path: Path
    relative_path: str
    environment: str
    sample_id: str
    condition: str
    filter_label: str
    filter_nm: float | None
    metrics: dict[str, float | str | int]


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    spectra = load_all_spectra()
    cameras = load_all_camera_measurements()
    references = build_reference_index(spectra)

    spectrum_features = [spectrum_feature_row(s, references) for s in spectra]
    camera_rows = [camera_row(c) for c in cameras]
    paired_rows = build_paired_rows(spectrum_features, camera_rows)
    reflectance_curves = build_reflectance_curves(spectra, references)

    write_csv(OUT_DIR / "spectra_features.csv", spectrum_features)
    write_csv(OUT_DIR / "camera_features.csv", camera_rows)
    write_csv(OUT_DIR / "camera_spectra_pairs.csv", paired_rows)
    write_long_spectra_csv(OUT_DIR / "spectra_long_sampled.csv", spectra, reflectance_curves)
    write_json_summary(OUT_DIR / "analysis_summary.json", spectra, cameras, spectrum_features, paired_rows)

    plot_spectra_overview(spectra, OUT_DIR / "raw_spectra_overview.png")
    plot_reflectance_overview(reflectance_curves, OUT_DIR / "reflectance_overview.png")
    plot_feature_comparison(spectrum_features, OUT_DIR / "spectral_feature_comparison.png")
    plot_camera_vs_spectrum(paired_rows, OUT_DIR / "camera_vs_spectrum_scatter.png")
    make_camera_contact_sheet(cameras, OUT_DIR / "camera_contact_sheet.jpg")
    if Workbook is not None:
        build_workbook(
            OUT_DIR / "real_leaf_camera_spectrometer_analysis.xlsx",
            spectrum_features,
            camera_rows,
            paired_rows,
        )
    build_dashboard(
        OUT_DIR / "real_leaf_analysis_dashboard.html",
        spectrum_features,
        camera_rows,
        paired_rows,
    )
    print(f"Wrote analysis to {OUT_DIR}")


def load_all_spectra() -> list[Spectrum]:
    spectra: list[Spectrum] = []
    for path in sorted(DATA_ROOT.rglob("*.txt")):
        parsed = read_spectrum(path)
        if parsed is None:
            continue
        wavelengths, intensities = parsed
        metadata = infer_metadata(path)
        spectra.append(
            Spectrum(
                path=path,
                relative_path=str(path.relative_to(DATA_ROOT)),
                wavelengths=wavelengths,
                intensities=intensities,
                **metadata,
            )
        )
    return spectra


def load_all_camera_measurements() -> list[CameraMeasurement]:
    measurements: list[CameraMeasurement] = []
    for path in sorted(DATA_ROOT.rglob("*.jpg")):
        metadata = infer_metadata(path)
        measurements.append(
            CameraMeasurement(
                path=path,
                relative_path=str(path.relative_to(DATA_ROOT)),
                environment=str(metadata["environment"]),
                sample_id=str(metadata["sample_id"]),
                condition=str(metadata["condition"]),
                filter_label=str(metadata["filter_label"]),
                filter_nm=metadata["filter_nm"] if isinstance(metadata["filter_nm"], float) else None,
                metrics=analyze_camera_image(path),
            )
        )
    return measurements


def read_spectrum(path: Path) -> tuple[np.ndarray, np.ndarray] | None:
    lines = path.read_text(encoding="utf-8", errors="ignore").splitlines()
    start = 0
    for index, line in enumerate(lines):
        if "Begin Spectral Data" in line:
            start = index + 1
            break
    pairs: list[tuple[float, float]] = []
    for line in lines[start:]:
        nums = re.findall(r"[-+]?(?:\d*\.\d+|\d+)(?:[eE][-+]?\d+)?", line)
        if len(nums) >= 2:
            pairs.append((float(nums[0]), float(nums[1])))
    if len(pairs) < 10:
        return None
    arr = np.asarray(pairs, dtype=float)
    return arr[:, 0], arr[:, 1]


def infer_metadata(path: Path) -> dict[str, object]:
    rel_parts = path.relative_to(DATA_ROOT).parts
    text = str(path.relative_to(DATA_ROOT)).lower()
    environment = rel_parts[0] if rel_parts else "unknown"
    name = path.stem.lower()

    filter_nm = infer_filter_nm(path.name)
    filter_label = "none" if filter_nm is None and ("fara" in name or "filtru" in name) else (
        str(int(filter_nm)) if filter_nm is not None else ""
    )

    if "referinta alba" in text or "refalb" in text:
        role = "white_reference"
        condition = "white_reference"
        sample_id = f"{environment}/white_reference"
    elif "refirinta neagra" in text or "referinta neagra" in text or "refnegru" in text:
        role = "dark_reference"
        condition = "dark_reference"
        sample_id = f"{environment}/dark_reference"
    elif "soare" in text or "sursa" in text:
        role = "source_light"
        condition = "source_light"
        sample_id = f"{environment}/source_light"
    else:
        role = "sample"
        if "nesan" in text:
            condition = "unhealthy"
        elif "galbena" in text:
            condition = "healthy_yellow"
        elif "frunza verde" in text or "\\san\\" in text or "sanatoasa" in text:
            condition = "healthy"
        else:
            condition = "sample_unknown"
        sample_id = infer_sample_id(rel_parts)

    modality = "camera" if path.suffix.lower() in {".jpg", ".jpeg", ".png"} else "spectrum"
    return {
        "environment": environment,
        "sample_id": sample_id,
        "condition": condition,
        "role": role,
        "modality": modality,
        "filter_label": filter_label,
        "filter_nm": filter_nm,
    }


def infer_sample_id(parts: tuple[str, ...]) -> str:
    if not parts:
        return "unknown"
    if parts[0].lower() == "afara":
        if len(parts) >= 3:
            return "/".join(parts[:3])
        return "/".join(parts[:-1])
    if parts[0].lower() == "in cutie":
        if len(parts) >= 4 and parts[-2].lower() in {"camera", "spectru"}:
            return "/".join(parts[:-2])
        return "/".join(parts[:-1])
    return "/".join(parts[:-1])


def infer_filter_nm(name: str) -> float | None:
    match = re.search(r"(532|556|580|680|725|850|940|950)", name)
    if not match:
        return None
    return float(match.group(1))


def analyze_camera_image(path: Path) -> dict[str, float | str | int]:
    image = Image.open(path).convert("RGB")
    rgb = np.asarray(image, dtype=np.uint8)
    mask = leaf_mask(rgb)
    pixels = rgb[mask]
    if pixels.size == 0:
        pixels = rgb.reshape(-1, 3)

    arr = pixels.astype(np.float32) / 255.0
    r, g, b = arr[:, 0], arr[:, 1], arr[:, 2]
    if cv2 is not None:
        hsv = cv2.cvtColor(pixels.reshape(-1, 1, 3), cv2.COLOR_RGB2HSV).reshape(-1, 3)
        hue = circular_mean_degrees(hsv[:, 0].astype(np.float32) * 2.0)
        sat = hsv[:, 1].astype(np.float32) / 255.0
        val = hsv[:, 2].astype(np.float32) / 255.0
    else:
        maxc = np.max(arr, axis=1)
        minc = np.min(arr, axis=1)
        sat = np.divide(maxc - minc, np.maximum(maxc, 1e-8))
        val = maxc
        hue = 0.0
    gcc = np.mean(g / np.maximum(r + g + b, 1e-8))
    exg = np.mean(2.0 * g - r - b)
    vari = np.nanmean((g - r) / np.where(np.abs(g + r - b) < 1e-8, np.nan, g + r - b))
    return {
        "width": image.width,
        "height": image.height,
        "leaf_pixels": int(mask.sum()),
        "leaf_area_fraction": round_float(mask.mean()),
        "mean_r": round_float(r.mean()),
        "mean_g": round_float(g.mean()),
        "mean_b": round_float(b.mean()),
        "brightness": round_float(val.mean()),
        "saturation": round_float(sat.mean()),
        "hue_deg": round_float(hue),
        "gcc": round_float(gcc),
        "exg": round_float(exg),
        "vari": round_float(vari),
    }


def leaf_mask(rgb: np.ndarray) -> np.ndarray:
    if cv2 is None:
        gray = np.asarray(Image.fromarray(rgb).convert("L"), dtype=np.uint8)
        mask = gray > np.percentile(gray, 65)
        return mask
    gray = cv2.cvtColor(rgb, cv2.COLOR_RGB2GRAY)
    blur = cv2.GaussianBlur(gray, (9, 9), 0)
    _, thresh = cv2.threshold(blur, 0, 255, cv2.THRESH_BINARY + cv2.THRESH_OTSU)
    mask = thresh > 0
    border = np.concatenate([mask[:20, :].ravel(), mask[-20:, :].ravel(), mask[:, :20].ravel(), mask[:, -20:].ravel()])
    if border.mean() > 0.5:
        mask = ~mask
    contours, _ = cv2.findContours(mask.astype(np.uint8) * 255, cv2.RETR_EXTERNAL, cv2.CHAIN_APPROX_SIMPLE)
    if contours:
        filled = np.zeros(mask.shape, dtype=np.uint8)
        cv2.drawContours(filled, [max(contours, key=cv2.contourArea)], -1, 255, cv2.FILLED)
        mask = filled > 0
    if mask.mean() < 0.02 or mask.mean() > 0.9:
        mask = gray > np.percentile(gray, 65)
    return mask


def build_reference_index(spectra: list[Spectrum]) -> dict[tuple[str, str, str], Spectrum]:
    refs: dict[tuple[str, str, str], Spectrum] = {}
    for spectrum in spectra:
        if spectrum.role not in {"white_reference", "dark_reference"}:
            continue
        keys = [(spectrum.environment, spectrum.role, spectrum.filter_label or "none")]
        if spectrum.filter_label != "none":
            keys.append((spectrum.environment, spectrum.role, "any"))
        for key in keys:
            refs.setdefault(key, spectrum)
    return refs


def choose_reference(
    spectrum: Spectrum,
    references: dict[tuple[str, str, str], Spectrum],
    role: str,
) -> Spectrum | None:
    candidates = [
        (spectrum.environment, role, spectrum.filter_label or "none"),
        (spectrum.environment, role, "none"),
        (spectrum.environment, role, "any"),
    ]
    for key in candidates:
        if key in references:
            return references[key]
    return None


def calibrated_reflectance(
    spectrum: Spectrum,
    references: dict[tuple[str, str, str], Spectrum],
) -> tuple[np.ndarray | None, str]:
    if spectrum.role != "sample":
        return None, "not_sample"
    white = choose_reference(spectrum, references, "white_reference")
    dark = choose_reference(spectrum, references, "dark_reference")
    if white is None or dark is None:
        return None, "missing_reference"
    if not (
        np.array_equal(spectrum.wavelengths, white.wavelengths)
        and np.array_equal(spectrum.wavelengths, dark.wavelengths)
    ):
        return None, "wavelength_grid_mismatch"
    denominator = white.intensities - dark.intensities
    valid = np.abs(denominator) > 1e-9
    result = np.full_like(spectrum.intensities, np.nan, dtype=float)
    np.divide(spectrum.intensities - dark.intensities, denominator, out=result, where=valid)
    return result, "ok"


def spectrum_feature_row(
    spectrum: Spectrum,
    references: dict[tuple[str, str, str], Spectrum],
) -> dict[str, object]:
    reflectance, status = calibrated_reflectance(spectrum, references)
    values = spectrum.intensities
    row: dict[str, object] = {
        "relative_path": spectrum.relative_path,
        "file": spectrum.path.name,
        "environment": spectrum.environment,
        "sample_id": spectrum.sample_id,
        "condition": spectrum.condition,
        "role": spectrum.role,
        "filter_label": spectrum.filter_label,
        "filter_nm": spectrum.filter_nm if spectrum.filter_nm is not None else "",
        "points": int(spectrum.wavelengths.size),
        "wavelength_min": round_float(np.nanmin(spectrum.wavelengths)),
        "wavelength_max": round_float(np.nanmax(spectrum.wavelengths)),
        "raw_min": round_float(np.nanmin(values)),
        "raw_max": round_float(np.nanmax(values)),
        "raw_mean": round_float(np.nanmean(values)),
        "raw_peak_wavelength_nm": round_float(spectrum.wavelengths[np.nanargmax(values)]),
        "raw_peak_intensity": round_float(np.nanmax(values)),
        "reflectance_status": status,
    }
    for wavelength in TARGET_WAVELENGTHS:
        row[f"raw_i_{wavelength}"] = round_float(np.interp(wavelength, spectrum.wavelengths, values))
        row[f"refl_{wavelength}"] = (
            round_float(np.interp(wavelength, spectrum.wavelengths, reflectance))
            if reflectance is not None
            else ""
        )
    row["raw_ratio_850_680"] = safe_ratio(row["raw_i_850"], row["raw_i_680"])
    row["raw_ratio_725_680"] = safe_ratio(row["raw_i_725"], row["raw_i_680"])
    if reflectance is not None:
        row["refl_ratio_850_680"] = safe_ratio(row["refl_850"], row["refl_680"])
        row["refl_ratio_725_680"] = safe_ratio(row["refl_725"], row["refl_680"])
    else:
        row["refl_ratio_850_680"] = ""
        row["refl_ratio_725_680"] = ""
    return row


def camera_row(camera: CameraMeasurement) -> dict[str, object]:
    return {
        "relative_path": camera.relative_path,
        "file": camera.path.name,
        "environment": camera.environment,
        "sample_id": camera.sample_id,
        "condition": camera.condition,
        "filter_label": camera.filter_label,
        "filter_nm": camera.filter_nm if camera.filter_nm is not None else "",
        **camera.metrics,
    }


def build_paired_rows(
    spectrum_rows: list[dict[str, object]],
    camera_rows: list[dict[str, object]],
) -> list[dict[str, object]]:
    sample_spectra = [
        row
        for row in spectrum_rows
        if row["role"] == "sample" and row["filter_label"] not in {"", "none"}
    ]
    cam_index = {
        (row["sample_id"], str(row["filter_label"])): row
        for row in camera_rows
        if str(row.get("filter_label", "")) not in {"", "none"}
    }
    paired: list[dict[str, object]] = []
    for spec in sample_spectra:
        key = (spec["sample_id"], str(spec["filter_label"]))
        cam = cam_index.get(key)
        if cam is None:
            continue
        wavelength_key = str(int(float(spec["filter_label"])))
        paired.append(
            {
                "environment": spec["environment"],
                "sample_id": spec["sample_id"],
                "condition": spec["condition"],
                "filter_label": spec["filter_label"],
                "filter_nm": spec["filter_nm"],
                "camera_file": cam["file"],
                "spectrum_file": spec["file"],
                "camera_brightness": cam["brightness"],
                "camera_saturation": cam["saturation"],
                "camera_gcc": cam["gcc"],
                "camera_exg": cam["exg"],
                "camera_vari": cam["vari"],
                "raw_intensity_at_filter": spec.get(f"raw_i_{wavelength_key}", ""),
                "reflectance_at_filter": spec.get(f"refl_{wavelength_key}", ""),
                "raw_ratio_850_680": spec.get("raw_ratio_850_680", ""),
                "refl_ratio_850_680": spec.get("refl_ratio_850_680", ""),
            }
        )
    return paired


def build_reflectance_curves(
    spectra: list[Spectrum],
    references: dict[tuple[str, str, str], Spectrum],
) -> dict[str, tuple[np.ndarray, np.ndarray]]:
    curves: dict[str, tuple[np.ndarray, np.ndarray]] = {}
    for spectrum in spectra:
        reflectance, status = calibrated_reflectance(spectrum, references)
        if status == "ok" and reflectance is not None:
            curves[spectrum.relative_path] = (spectrum.wavelengths, reflectance)
    return curves


def plot_spectra_overview(spectra: list[Spectrum], output: Path) -> None:
    fig, axes = plt.subplots(1, 2, figsize=(13, 5), sharex=True)
    for ax, environment in zip(axes, ["afara", "in cutie"]):
        for spectrum in spectra:
            if spectrum.environment != environment or spectrum.role != "sample" or spectrum.filter_label == "none":
                continue
            label = f"{spectrum.condition} {spectrum.filter_label}"
            ax.plot(spectrum.wavelengths, spectrum.intensities, linewidth=0.9, alpha=0.55, label=label)
        ax.set_title(f"Raw spectra: {environment}")
        ax.set_xlabel("Wavelength (nm)")
        ax.set_ylabel("Intensity")
        ax.grid(alpha=0.25)
    handles, labels = axes[1].get_legend_handles_labels()
    if handles:
        axes[1].legend(handles[:12], labels[:12], fontsize=7)
    fig.tight_layout()
    fig.savefig(output, dpi=170)
    plt.close(fig)


def plot_reflectance_overview(curves: dict[str, tuple[np.ndarray, np.ndarray]], output: Path) -> None:
    fig, ax = plt.subplots(figsize=(11, 5.8))
    for relative_path, (wavelengths, reflectance) in curves.items():
        if "Referinta" in relative_path or "ref" in relative_path.lower():
            continue
        label = relative_path.replace("\\spectru\\", "\\").replace("in cutie\\", "").replace("afara\\", "")
        ax.plot(wavelengths, np.clip(reflectance, -0.2, 2.5), linewidth=1.0, alpha=0.65, label=label)
    ax.set_title("Calibrated reflectance curves where references are available")
    ax.set_xlabel("Wavelength (nm)")
    ax.set_ylabel("Reflectance (clipped to -0.2..2.5 for display)")
    ax.grid(alpha=0.25)
    ax.legend(fontsize=6, ncol=2)
    fig.tight_layout()
    fig.savefig(output, dpi=170)
    plt.close(fig)


def plot_feature_comparison(rows: list[dict[str, object]], output: Path) -> None:
    samples = [row for row in rows if row["role"] == "sample" and row["filter_label"] not in {"", "none"}]
    conditions = sorted({str(row["condition"]) for row in samples})
    wavelengths = [532, 556, 680, 725, 850, 940]
    fig, axes = plt.subplots(1, 2, figsize=(12, 5))
    for condition in conditions:
        xs: list[int] = []
        raw_means: list[float] = []
        refl_means: list[float] = []
        for wavelength in wavelengths:
            subset = [row for row in samples if row["condition"] == condition and str(row["filter_label"]) == str(wavelength)]
            if not subset:
                continue
            xs.append(wavelength)
            raw_means.append(float(np.nanmean([float(row[f"raw_i_{wavelength}"]) for row in subset])))
            refl_values = [to_float(row[f"refl_{wavelength}"]) for row in subset if row[f"refl_{wavelength}"] != ""]
            refl_means.append(float(np.nanmean(refl_values)) if refl_values else np.nan)
        axes[0].plot(xs, raw_means, marker="o", label=condition)
        axes[1].plot(xs, refl_means, marker="o", label=condition)
    axes[0].set_title("Mean raw intensity at matching filter")
    axes[1].set_title("Mean calibrated reflectance at matching filter")
    for ax in axes:
        ax.set_xlabel("Filter wavelength (nm)")
        ax.grid(alpha=0.25)
        ax.legend(fontsize=8)
    axes[0].set_ylabel("Raw intensity")
    axes[1].set_ylabel("Reflectance")
    fig.tight_layout()
    fig.savefig(output, dpi=170)
    plt.close(fig)


def plot_camera_vs_spectrum(rows: list[dict[str, object]], output: Path) -> None:
    valid = [
        row
        for row in rows
        if row.get("reflectance_at_filter") not in {"", None}
        and is_finite(row.get("camera_brightness"))
        and is_finite(row.get("reflectance_at_filter"))
    ]
    fig, ax = plt.subplots(figsize=(8, 5.5))
    colors = {"healthy": "#267a3f", "unhealthy": "#b24835", "healthy_yellow": "#b8871f"}
    for row in valid:
        ax.scatter(
            float(row["camera_brightness"]),
            float(row["reflectance_at_filter"]),
            color=colors.get(str(row["condition"]), "#444"),
            s=48,
            alpha=0.8,
        )
        ax.text(float(row["camera_brightness"]), float(row["reflectance_at_filter"]), str(row["filter_label"]), fontsize=7)
    ax.set_title("Camera brightness vs calibrated spectrum at matching filter")
    ax.set_xlabel("Camera leaf brightness")
    ax.set_ylabel("Calibrated reflectance at same wavelength")
    ax.grid(alpha=0.25)
    fig.tight_layout()
    fig.savefig(output, dpi=170)
    plt.close(fig)


def make_camera_contact_sheet(cameras: list[CameraMeasurement], output: Path) -> None:
    tiles = []
    for camera in cameras:
        image = Image.open(camera.path).convert("RGB")
        image.thumbnail((220, 124))
        canvas = Image.new("RGB", (240, 175), "white")
        canvas.paste(image, (10, 10))
        draw = ImageDraw.Draw(canvas)
        label = f"{camera.sample_id.split('/')[-1]}\n{camera.filter_label or camera.path.stem} {camera.condition}"
        draw.text((10, 140), label[:60], fill="black")
        tiles.append(canvas)
    cols = 5
    sheet = Image.new("RGB", (cols * 240, math.ceil(len(tiles) / cols) * 175), (238, 240, 235))
    for idx, tile in enumerate(tiles):
        sheet.paste(tile, ((idx % cols) * 240, (idx // cols) * 175))
    sheet.save(output, quality=92)


def write_long_spectra_csv(
    path: Path,
    spectra: list[Spectrum],
    reflectance_curves: dict[str, tuple[np.ndarray, np.ndarray]],
) -> None:
    # Keep this sampled enough to inspect in Excel without making a giant file.
    rows = []
    for spectrum in spectra:
        step = max(int(spectrum.wavelengths.size // 256), 1)
        reflectance = reflectance_curves.get(spectrum.relative_path, (None, None))[1]
        for idx in range(0, spectrum.wavelengths.size, step):
            rows.append(
                {
                    "relative_path": spectrum.relative_path,
                    "environment": spectrum.environment,
                    "sample_id": spectrum.sample_id,
                    "condition": spectrum.condition,
                    "role": spectrum.role,
                    "filter_label": spectrum.filter_label,
                    "wavelength_nm": round_float(spectrum.wavelengths[idx]),
                    "raw_intensity": round_float(spectrum.intensities[idx]),
                    "reflectance": round_float(reflectance[idx]) if reflectance is not None else "",
                }
            )
    write_csv(path, rows)


def write_json_summary(
    path: Path,
    spectra: list[Spectrum],
    cameras: list[CameraMeasurement],
    spectrum_features: list[dict[str, object]],
    paired_rows: list[dict[str, object]],
) -> None:
    sample_rows = [row for row in spectrum_features if row["role"] == "sample"]
    summary = {
        "data_root": str(DATA_ROOT),
        "spectra_count": len(spectra),
        "camera_count": len(cameras),
        "paired_camera_spectrum_rows": len(paired_rows),
        "reflectance_ok_count": sum(1 for row in sample_rows if row["reflectance_status"] == "ok"),
        "conditions": sorted({row["condition"] for row in sample_rows}),
        "target_wavelengths": list(TARGET_WAVELENGTHS),
    }
    path.write_text(json.dumps(summary, indent=2), encoding="utf-8")


def build_workbook(
    path: Path,
    spectrum_features: list[dict[str, object]],
    camera_rows: list[dict[str, object]],
    paired_rows: list[dict[str, object]],
) -> None:
    wb = Workbook()
    wb.remove(wb.active)
    add_sheet(wb, "Summary", summary_rows(spectrum_features, camera_rows, paired_rows))
    add_sheet(wb, "Spectra Features", spectrum_features)
    add_sheet(wb, "Camera Features", camera_rows)
    add_sheet(wb, "Camera Spectrum Pairs", paired_rows)
    ws = wb["Summary"]
    if ws.max_row >= 4:
        chart = LineChart()
        chart.title = "Dataset Counts"
        chart.y_axis.title = "Count"
        data = Reference(ws, min_col=2, max_col=2, min_row=1, max_row=min(ws.max_row, 5))
        cats = Reference(ws, min_col=1, min_row=2, max_row=min(ws.max_row, 5))
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(cats)
        ws.add_chart(chart, "E2")
    format_workbook(wb)
    wb.save(path)


def summary_rows(
    spectrum_features: list[dict[str, object]],
    camera_rows: list[dict[str, object]],
    paired_rows: list[dict[str, object]],
) -> list[dict[str, object]]:
    sample_rows = [row for row in spectrum_features if row["role"] == "sample"]
    rows: list[dict[str, object]] = [
        {"metric": "spectra_files_parsed", "value": len(spectrum_features), "notes": "TXT spectra with 2048 wavelength/intensity points"},
        {"metric": "sample_spectra", "value": len(sample_rows), "notes": "Excludes references/source lights"},
        {"metric": "camera_images", "value": len(camera_rows), "notes": "JPG files analyzed"},
        {"metric": "paired_camera_spectrum_rows", "value": len(paired_rows), "notes": "Same sample_id and wavelength/filter"},
        {"metric": "reflectance_calibrated_samples", "value": sum(1 for row in sample_rows if row["reflectance_status"] == "ok"), "notes": "Reflectance = (sample - dark) / (white - dark)"},
    ]
    for condition in sorted({row["condition"] for row in sample_rows}):
        subset = [row for row in sample_rows if row["condition"] == condition]
        rows.append({"metric": f"{condition}_sample_spectra", "value": len(subset), "notes": ""})
    return rows


def build_dashboard(
    path: Path,
    spectrum_features: list[dict[str, object]],
    camera_rows: list[dict[str, object]],
    paired_rows: list[dict[str, object]],
) -> None:
    summary = summary_rows(spectrum_features, camera_rows, paired_rows)
    payload = json.dumps({"summary": summary, "pairs": paired_rows[:200]}, ensure_ascii=False)
    html = f"""<!doctype html>
<html lang="en">
<head>
  <meta charset="utf-8" />
  <meta name="viewport" content="width=device-width, initial-scale=1" />
  <title>Real Leaf Spectral Analysis</title>
  <style>
    body {{ margin:0; font-family: Segoe UI, Arial, sans-serif; background:#f6f7f3; color:#17211b; }}
    header {{ padding:28px 32px; background:white; border-bottom:1px solid #d9ded5; }}
    main {{ padding:24px 32px 40px; max-width:1280px; margin:auto; }}
    h1 {{ margin:0 0 8px; font-size:28px; }}
    h2 {{ font-size:18px; margin:0 0 12px; }}
    p {{ color:#5f6b62; }}
    .grid {{ display:grid; grid-template-columns:repeat(4,1fr); gap:14px; }}
    .panel {{ background:white; border:1px solid #d9ded5; border-radius:8px; padding:16px; margin-bottom:16px; }}
    .kpi {{ font-size:30px; font-weight:750; }}
    .label {{ color:#667168; font-size:12px; text-transform:uppercase; }}
    img {{ width:100%; border:1px solid #d9ded5; border-radius:8px; background:white; }}
    table {{ width:100%; border-collapse:collapse; font-size:13px; }}
    th,td {{ padding:8px; border-bottom:1px solid #e2e6df; text-align:left; }}
    th {{ color:#667168; }}
    @media (max-width:900px) {{ .grid {{ grid-template-columns:1fr; }} main,header {{ padding-left:18px; padding-right:18px; }} }}
  </style>
</head>
<body>
  <header>
    <h1>Real Leaf Camera + Spectrometer Analysis</h1>
    <p>Now using exported TXT spectra with 2048 wavelength-intensity points, camera JPGs, and white/dark reference calibration where available.</p>
  </header>
  <main>
    <section class="grid" id="kpis"></section>
    <section class="panel">
      <h2>Spectral Graphs</h2>
      <div class="grid" style="grid-template-columns:1fr 1fr">
        <img src="raw_spectra_overview.png" alt="Raw spectra overview" />
        <img src="reflectance_overview.png" alt="Reflectance overview" />
      </div>
    </section>
    <section class="panel">
      <h2>Feature Comparison</h2>
      <div class="grid" style="grid-template-columns:1fr 1fr">
        <img src="spectral_feature_comparison.png" alt="Spectral feature comparison" />
        <img src="camera_vs_spectrum_scatter.png" alt="Camera vs spectrum scatter" />
      </div>
    </section>
    <section class="panel">
      <h2>Camera Measurements</h2>
      <img src="camera_contact_sheet.jpg" alt="Camera contact sheet" />
    </section>
    <section class="panel">
      <h2>First Paired Rows</h2>
      <table id="pairs"></table>
    </section>
  </main>
  <script>
    const DATA = {payload};
    const wanted = ["spectra_files_parsed","sample_spectra","camera_images","paired_camera_spectrum_rows"];
    document.getElementById("kpis").innerHTML = DATA.summary.filter(r => wanted.includes(r.metric)).map(r => `<div class="panel"><div class="label">${{r.metric}}</div><div class="kpi">${{r.value}}</div><p>${{r.notes}}</p></div>`).join("");
    const pairs = DATA.pairs.slice(0, 20);
    const cols = ["sample_id","condition","filter_label","camera_brightness","camera_gcc","reflectance_at_filter","raw_intensity_at_filter"];
    document.getElementById("pairs").innerHTML = `<thead><tr>${{cols.map(c=>`<th>${{c}}</th>`).join("")}}</tr></thead><tbody>${{pairs.map(r=>`<tr>${{cols.map(c=>`<td>${{r[c] ?? ""}}</td>`).join("")}}</tr>`).join("")}}</tbody>`;
  </script>
</body>
</html>
"""
    path.write_text(html, encoding="utf-8")


def add_sheet(wb: Workbook, name: str, rows: list[dict[str, object]]) -> None:
    ws = wb.create_sheet(name)
    if not rows:
        ws.append(["No data"])
        return
    headers = list(rows[0].keys())
    ws.append(headers)
    for row in rows:
        ws.append([row.get(header, "") for header in headers])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def format_workbook(wb: Workbook) -> None:
    fill = PatternFill("solid", fgColor="1F4E78")
    font = Font(color="FFFFFF", bold=True)
    for ws in wb.worksheets:
        for cell in ws[1]:
            cell.fill = fill
            cell.font = font
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for col in ws.columns:
            letter = get_column_letter(col[0].column)
            width = min(max(max(len(str(cell.value)) if cell.value is not None else 0 for cell in col[:200]) + 2, 10), 48)
            ws.column_dimensions[letter].width = width


def write_csv(path: Path, rows: list[dict[str, object]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    if not rows:
        path.write_text("", encoding="utf-8")
        return
    headers: list[str] = []
    for row in rows:
        for key in row:
            if key not in headers:
                headers.append(key)
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=headers)
        writer.writeheader()
        writer.writerows(rows)


def circular_mean_degrees(values: np.ndarray) -> float:
    radians = np.deg2rad(values.astype(float))
    return float((np.rad2deg(np.arctan2(np.mean(np.sin(radians)), np.mean(np.cos(radians)))) + 360.0) % 360.0)


def safe_ratio(num: object, den: object) -> float | str:
    n, d = to_float(num), to_float(den)
    if not np.isfinite(n) or not np.isfinite(d) or abs(d) < 1e-12:
        return ""
    return round_float(n / d)


def to_float(value: object) -> float:
    try:
        return float(value)
    except (TypeError, ValueError):
        return float("nan")


def is_finite(value: object) -> bool:
    return np.isfinite(to_float(value))


def round_float(value: object, digits: int = 6) -> float:
    number = to_float(value)
    return round(number, digits) if np.isfinite(number) else float("nan")


if __name__ == "__main__":
    main()
