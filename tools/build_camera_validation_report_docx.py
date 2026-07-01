"""Build the camera validation report DOCX."""

from __future__ import annotations

import csv
import json
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"C:\Users\david\OneDrive\Desktop\Archive\Projects\CodeX")
REPORT_DIR = ROOT / "outputs" / "camera_validation_report"
GRAPH_DIR = REPORT_DIR / "graphs"
DOCX_PATH = REPORT_DIR / "camera_validation_report.docx"


def main() -> None:
    summary = json.loads((REPORT_DIR / "validation_summary.json").read_text(encoding="utf-8"))
    corr_rows = read_csv(REPORT_DIR / "per_sample_camera_spectrum_correlations.csv")
    graph_rows = read_csv(REPORT_DIR / "graph_manifest.csv")

    doc = Document()
    setup_document(doc)
    add_title_page(doc, summary)
    add_executive_summary(doc, summary, corr_rows)
    add_method(doc)
    add_key_results(doc, summary, corr_rows)
    add_figures(doc)
    add_app_use(doc)
    add_limitations(doc)
    add_graph_appendix(doc, graph_rows)
    doc.save(DOCX_PATH)
    print(DOCX_PATH)


def setup_document(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    styles = doc.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(10.5)
    styles["Heading 1"].font.name = "Aptos Display"
    styles["Heading 1"].font.size = Pt(18)
    styles["Heading 1"].font.color.rgb = RGBColor(31, 78, 121)
    styles["Heading 2"].font.name = "Aptos"
    styles["Heading 2"].font.size = Pt(14)
    styles["Heading 2"].font.color.rgb = RGBColor(47, 98, 62)


def add_title_page(doc: Document, summary: dict) -> None:
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Camera Validation Report\n")
    run.bold = True
    run.font.size = Pt(26)
    run.font.color.rgb = RGBColor(31, 78, 121)
    subtitle = title.add_run("Full-spectrum spectrometer comparison with filtered camera measurements")
    subtitle.font.size = Pt(13)
    subtitle.font.color.rgb = RGBColor(90, 100, 95)

    doc.add_paragraph()
    add_callout(
        doc,
        "Main conclusion",
        (
            "The camera produces structured, wavelength-dependent data and tracks the spectrometer "
            "well in several matched sample series. The current dataset supports a credible proof of "
            "concept, especially with per-sample normalized fingerprints. It is not yet a calibration-grade "
            "proof across every sample because exposure/reference control appears inconsistent."
        ),
    )
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    hdr[0].text = "Measurement"
    hdr[1].text = "Value"
    rows = [
        ("Paired camera-spectrum measurements", summary["paired_measurements"]),
        ("Per-sample fingerprint comparisons", summary["per_sample_correlations"]),
        ("Median camera/spectrometer fingerprint correlation", f"{summary['median_per_sample_normalized_correlation']:.3f}"),
        ("Full-spectrum classifier accuracy", f"{summary['full_spectrum_classifier']['accuracy']:.1%}"),
        ("PCA variance explained by PC1 + PC2", f"{summary['pca']['pc1_explained_variance'] + summary['pca']['pc2_explained_variance']:.1%}"),
    ]
    for label, value in rows:
        cells = table.add_row().cells
        cells[0].text = str(label)
        cells[1].text = str(value)
    doc.add_page_break()


def add_executive_summary(doc: Document, summary: dict, corr_rows: list[dict[str, str]]) -> None:
    doc.add_heading("Executive Summary", level=1)
    bullets = [
        "The new TXT files contain real spectrometer data: each spectrum has 2048 wavelength-intensity points across the full measured range.",
        "The analysis intentionally uses the complete spectrum for full-spectrum area, normalized curve shape, centroid, PCA, and similarity calculations.",
        "The strongest camera-vs-spectrometer evidence is the per-sample fingerprint comparison: several sample series show moderate to strong tracking between camera response and spectrometer full-spectrum response.",
        "The median per-sample normalized camera/spectrometer correlation is "
        f"{summary['median_per_sample_normalized_correlation']:.3f}.",
        "The global raw correlation is weak, which is expected if camera exposure, white balance, sample placement, or reference conditions changed between captures.",
        "A naive full-spectrum healthy/unhealthy classifier is weak in this dataset, so the report does not claim disease classification is solved. The defensible claim is that the camera captures useful wavelength-dependent signal that can be aligned to spectrometer measurements under controlled conditions.",
    ]
    for item in bullets:
        doc.add_paragraph(item, style="List Bullet")

    doc.add_heading("Most Useful Evidence For Your App", level=2)
    table = doc.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    headers = ["Sample", "Condition", "Filters", "Camera vs spectrum r", "Interpretation"]
    for idx, text in enumerate(headers):
        table.rows[0].cells[idx].text = text
    for row in corr_rows:
        cells = table.add_row().cells
        cells[0].text = row["sample_id"].replace("in cutie/", "")
        cells[1].text = row["condition"]
        cells[2].text = row["paired_filters"]
        cells[3].text = row["pearson_normalized_fingerprints"]
        r = safe_float(row["pearson_normalized_fingerprints"])
        cells[4].text = "strong match" if r >= 0.7 else "moderate match" if r >= 0.5 else "weak or inconsistent"


def add_method(doc: Document) -> None:
    doc.add_heading("Method", level=1)
    doc.add_paragraph(
        "The analysis parsed all available spectrometer TXT files and camera JPG files from the real data folder. "
        "For every matched camera/spectrum pair, the matching key was sample folder plus filter label. The spectra "
        "were smoothed with a Savitzky-Golay filter, but no wavelength range was removed; the full measured curve "
        "was kept for full-spectrum features."
    )
    method_rows = [
        ("Full-spectrum area", "Integral of the smoothed, dark-corrected positive signal over the entire measured wavelength range."),
        ("Normalized spectrum", "Z-normalized smoothed curve using all 2048 points."),
        ("Camera brightness", "Mean leaf-pixel brightness from the segmented camera image."),
        ("Fingerprint correlation", "Correlation across filters between camera response and spectrometer full-spectrum area within the same sample."),
        ("PCA", "Dimensionality reduction of the complete normalized spectrum, not selected bands."),
    ]
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.rows[0].cells[0].text = "Feature"
    table.rows[0].cells[1].text = "Meaning"
    for a, b in method_rows:
        cells = table.add_row().cells
        cells[0].text = a
        cells[1].text = b


def add_key_results(doc: Document, summary: dict, corr_rows: list[dict[str, str]]) -> None:
    doc.add_heading("Findings", level=1)
    doc.add_heading("1. The camera signal is structured, not random", level=2)
    doc.add_paragraph(
        "The filtered camera images produce repeatable changes across filter labels. This is visible in the "
        "fingerprint plots and camera feature heatmap. The signal is not merely noise; it changes with filter, "
        "sample, and condition."
    )
    doc.add_heading("2. The spectrometer confirms strong wavelength-dependent structure", level=2)
    doc.add_paragraph(
        "The spectrometer curves show strong filter-dependent peaks and broad full-spectrum differences. This makes "
        "them suitable as a reference instrument for validating the camera response."
    )
    doc.add_heading("3. Per-sample comparison is the right proof strategy", level=2)
    doc.add_paragraph(
        "Raw global correlation is weak because the dataset combines different leaves, conditions, references, and likely "
        "camera exposure states. When each sample is normalized internally and compared as a response fingerprint, several "
        "samples show useful agreement."
    )
    for row in corr_rows:
        r = safe_float(row["pearson_normalized_fingerprints"])
        if r >= 0.6:
            doc.add_paragraph(
                f"{row['sample_id'].replace('in cutie/', '')}: r={r:.3f}, R^2={safe_float(row['normalized_r_squared']):.3f}.",
                style="List Bullet",
            )
    doc.add_heading("4. The present dataset does not prove universal calibration yet", level=2)
    doc.add_paragraph(
        "Some sample series are weak or inverted, and the full-spectrum classifier does not separate healthy/unhealthy reliably "
        f"(accuracy {summary['full_spectrum_classifier']['accuracy']:.1%}). This does not mean the camera is useless; it means "
        "the proof should be framed as a working, wavelength-sensitive camera prototype, not a finished calibrated instrument."
    )


def add_figures(doc: Document) -> None:
    doc.add_heading("Core Figures", level=1)
    figures = [
        ("03_condition_mean_full_spectrum.png", "Figure 1. Mean full-spectrum curves by condition."),
        ("04_healthy_minus_unhealthy_difference.png", "Figure 2. Full-spectrum difference curve: healthy minus unhealthy."),
        ("05_full_spectrum_pca.png", "Figure 3. PCA of complete normalized spectra."),
        ("07_camera_vs_spectrometer_normalized.png", "Figure 4. Within-sample normalized camera response versus spectrometer full-spectrum response."),
        ("09_per_sample_correlation_bars.png", "Figure 5. Per-sample correlation between camera and spectrometer fingerprints."),
        ("10_fingerprint_overlays_by_sample.png", "Figure 6. Camera and spectrometer response fingerprints by sample."),
        ("17_spectrum_area_by_condition_filter.png", "Figure 7. Spectrometer full-spectrum area by condition and filter."),
        ("18_full_spectrum_similarity_heatmap.png", "Figure 8. Full-spectrum similarity heatmap."),
    ]
    for filename, caption in figures:
        add_figure(doc, GRAPH_DIR / filename, caption)


def add_app_use(doc: Document) -> None:
    doc.add_heading("How To Use This In The App", level=1)
    doc.add_paragraph(
        "Use the camera as a multispectral response sensor by comparing response fingerprints rather than isolated raw "
        "brightness values. In the app, every sample should produce both a camera fingerprint and, when available, a "
        "spectrometer fingerprint."
    )
    steps = [
        "Segment the leaf in each filtered camera image.",
        "Compute camera brightness, saturation, GCC, and other image features.",
        "Normalize those features within the sample across filters.",
        "For spectrometer validation, compute full-spectrum area and normalized full-spectrum shape using all wavelengths.",
        "Report camera/spectrometer fingerprint correlation as a confidence score.",
        "Flag low-correlation samples for retake or calibration review.",
    ]
    for step in steps:
        doc.add_paragraph(step, style="List Number")
    add_callout(
        doc,
        "Best proof-of-concept claim",
        (
            "The camera captures wavelength-dependent leaf response patterns that can be compared to a spectrometer. "
            "Under controlled sample series, the camera fingerprint moderately to strongly tracks the spectrometer "
            "full-spectrum fingerprint."
        ),
    )


def add_limitations(doc: Document) -> None:
    doc.add_heading("Limitations And Next Measurements", level=1)
    limitations = [
        "The JPG files do not expose useful EXIF exposure metadata, so changing exposure/white balance cannot be corrected after the fact.",
        "White-reference correction helps some samples but hurts others, suggesting the reference setup was not fully matched to every sample.",
        "The camera and spectrometer likely do not observe exactly the same spatial region of the leaf.",
        "Some classes are confounded with filter and environment. A naive classifier should not be used as the headline proof.",
    ]
    for item in limitations:
        doc.add_paragraph(item, style="List Bullet")
    doc.add_heading("Recommended retake protocol", level=2)
    retake = [
        "Lock camera exposure, ISO, focus, and white balance.",
        "Capture white and dark references for every filter in the same session.",
        "Use the same leaf ROI for camera and spectrometer measurements.",
        "Take at least three repeats per condition/filter.",
        "Keep filenames machine-readable: sample_condition_filter_repeat.",
    ]
    for item in retake:
        doc.add_paragraph(item, style="List Bullet")


def add_graph_appendix(doc: Document, graph_rows: list[dict[str, str]]) -> None:
    doc.add_heading("Graph Appendix", level=1)
    doc.add_paragraph(
        "All candidate graphs were generated as separate PNG files in the graphs folder. "
        "Use this list to choose which ones to include in a final presentation or app documentation."
    )
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.rows[0].cells[0].text = "Graph file"
    table.rows[0].cells[1].text = "Title"
    for row in graph_rows:
        cells = table.add_row().cells
        cells[0].text = row["file"]
        cells[1].text = row["title"]


def add_callout(doc: Document, title: str, body: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    cell = table.rows[0].cells[0]
    p = cell.paragraphs[0]
    run = p.add_run(title + "\n")
    run.bold = True
    run.font.color.rgb = RGBColor(31, 78, 121)
    p.add_run(body)


def add_figure(doc: Document, path: Path, caption: str) -> None:
    if not path.exists():
        doc.add_paragraph(f"Missing figure: {path.name}")
        return
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    run.add_picture(str(path), width=Inches(6.7))
    cap = doc.add_paragraph(caption)
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.runs[0].italic = True


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open(newline="", encoding="utf-8") as handle:
        return list(csv.DictReader(handle))


def safe_float(value: str) -> float:
    try:
        return float(value)
    except (TypeError, ValueError):
        return float("nan")


if __name__ == "__main__":
    main()
