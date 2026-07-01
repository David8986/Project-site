"""Create a visual camera/spectrometer calibration scheme document."""

from __future__ import annotations

from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "outputs" / "calibration_scheme"
PNG_PATH = OUT_DIR / "schema_calibrare_camera_spectrometru.png"
DOCX_PATH = OUT_DIR / "schema_calibrare_camera_spectrometru.docx"


COLORS = {
    "paper": "#f7f7f3",
    "ink": "#1f2933",
    "muted": "#52616b",
    "blue": "#d9ebff",
    "blue_edge": "#2f6fa7",
    "green": "#dff3e4",
    "green_edge": "#2f855a",
    "amber": "#fff1cc",
    "amber_edge": "#b7791f",
    "red": "#ffe0df",
    "red_edge": "#b83232",
    "gray": "#eef1f4",
    "gray_edge": "#718096",
    "dark": "#18212f",
}


def _font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    candidates = [
        Path("C:/Windows/Fonts/segoeui.ttf"),
        Path("C:/Windows/Fonts/segoeuib.ttf") if bold else Path("C:/Windows/Fonts/segoeui.ttf"),
        Path("C:/Windows/Fonts/arialbd.ttf") if bold else Path("C:/Windows/Fonts/arial.ttf"),
    ]
    for path in candidates:
        if path.exists():
            return ImageFont.truetype(str(path), size=size)
    return ImageFont.load_default()


def _text_size(draw: ImageDraw.ImageDraw, text: str, font: ImageFont.ImageFont) -> tuple[int, int]:
    bbox = draw.multiline_textbbox((0, 0), text, font=font, spacing=6)
    return int(bbox[2] - bbox[0]), int(bbox[3] - bbox[1])


def _wrap(draw: ImageDraw.ImageDraw, text: str, font: ImageFont.ImageFont, width: int) -> str:
    words = text.split()
    lines: list[str] = []
    current = ""
    for word in words:
        candidate = word if not current else f"{current} {word}"
        if draw.textlength(candidate, font=font) <= width:
            current = candidate
            continue
        if current:
            lines.append(current)
        current = word
    if current:
        lines.append(current)
    return "\n".join(lines)


def _box(
    draw: ImageDraw.ImageDraw,
    xy: tuple[int, int, int, int],
    title: str,
    body: str,
    fill: str,
    outline: str,
) -> None:
    x0, y0, x1, y1 = xy
    draw.rounded_rectangle(xy, radius=22, fill=fill, outline=outline, width=4)
    title_font = _font(28, bold=True)
    body_font = _font(21)
    draw.text((x0 + 24, y0 + 18), title, font=title_font, fill=COLORS["ink"])
    wrapped = _wrap(draw, body, body_font, x1 - x0 - 48)
    draw.multiline_text((x0 + 24, y0 + 58), wrapped, font=body_font, fill=COLORS["ink"], spacing=7)


def _arrow(draw: ImageDraw.ImageDraw, start: tuple[int, int], end: tuple[int, int], color: str = "#2d3748") -> None:
    draw.line([start, end], fill=color, width=5)
    sx, sy = start
    ex, ey = end
    if abs(ex - sx) >= abs(ey - sy):
        direction = 1 if ex >= sx else -1
        points = [(ex, ey), (ex - 18 * direction, ey - 11), (ex - 18 * direction, ey + 11)]
    else:
        direction = 1 if ey >= sy else -1
        points = [(ex, ey), (ex - 11, ey - 18 * direction), (ex + 11, ey - 18 * direction)]
    draw.polygon(points, fill=color)


def _decision(
    draw: ImageDraw.ImageDraw,
    center: tuple[int, int],
    size: tuple[int, int],
    title: str,
    body: str,
) -> None:
    cx, cy = center
    w, h = size
    points = [(cx, cy - h // 2), (cx + w // 2, cy), (cx, cy + h // 2), (cx - w // 2, cy)]
    draw.polygon(points, fill=COLORS["amber"], outline=COLORS["amber_edge"])
    draw.line(points + [points[0]], fill=COLORS["amber_edge"], width=4)
    title_font = _font(23, bold=True)
    body_font = _font(18)
    wrapped_title = _wrap(draw, title, title_font, w - 90)
    tw, th = _text_size(draw, wrapped_title, title_font)
    draw.multiline_text((cx - tw // 2, cy - 42), wrapped_title, font=title_font, fill=COLORS["ink"], align="center", spacing=4)
    wrapped_body = _wrap(draw, body, body_font, w - 110)
    bw, _ = _text_size(draw, wrapped_body, body_font)
    draw.multiline_text((cx - bw // 2, cy + 12), wrapped_body, font=body_font, fill=COLORS["ink"], align="center", spacing=4)


def build_png() -> Path:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    image = Image.new("RGB", (1800, 1120), COLORS["paper"])
    draw = ImageDraw.Draw(image)
    title_font = _font(42, bold=True)
    subtitle_font = _font(24)
    small_font = _font(19)

    draw.text((70, 48), "Schema completa: calibrare camera -> spectrometru", font=title_font, fill=COLORS["dark"])
    draw.text(
        (72, 104),
        "Fluxul folosit de aplicatie pentru pozele cu filtre si compararea cu masuratorile spectrometrului",
        font=subtitle_font,
        fill=COLORS["muted"],
    )

    boxes = {
        "photos": (70, 175, 410, 340),
        "brightness": (510, 175, 850, 340),
        "mask": (950, 175, 1290, 340),
        "mean": (1390, 175, 1730, 340),
        "white": (70, 455, 410, 620),
        "spectro": (510, 455, 850, 620),
        "regression": (950, 455, 1290, 620),
        "output": (1390, 455, 1730, 620),
    }

    _box(draw, boxes["photos"], "1. Intrare", "Poze cu filtre: 532, 556, 680, 725, 850, 940 nm", COLORS["blue"], COLORS["blue_edge"])
    _box(draw, boxes["brightness"], "2. Pixel", "I_lambda(x,y) = max(R,G,B) / 255", COLORS["gray"], COLORS["gray_edge"])
    _box(draw, boxes["mask"], "3. Masca", "NDVI = (I850 - I680) / (I850 + I680 + eps)", COLORS["green"], COLORS["green_edge"])
    _box(draw, boxes["mean"], "4. Medie", "C_lambda = media pixelilor frunza din banda lambda", COLORS["gray"], COLORS["gray_edge"])
    _box(draw, boxes["white"], "5. Referinta alba", "N_lambda = C_lambda / W_lambda", COLORS["green"], COLORS["green_edge"])
    _box(draw, boxes["spectro"], "6. Tinta reala", "S_lambda = mediana spectrometrului in [lambda - 10, lambda + 10] nm", COLORS["blue"], COLORS["blue_edge"])
    _box(draw, boxes["regression"], "7. Model", "S_hat_lambda = b_lambda + suma beta_lambda,k * ((N_k - mu_k) / sigma_k)", COLORS["amber"], COLORS["amber_edge"])
    _box(draw, boxes["output"], "8. Rezultat", "Valori estimate de spectrometru + NDVI, NDRE, GNDVI, NDWI", COLORS["green"], COLORS["green_edge"])

    for left, right in [("photos", "brightness"), ("brightness", "mask"), ("mask", "mean")]:
        _arrow(draw, ((boxes[left][2]), (boxes[left][1] + boxes[left][3]) // 2), (boxes[right][0], (boxes[right][1] + boxes[right][3]) // 2))
    _arrow(draw, ((boxes["mean"][0] + boxes["mean"][2]) // 2, boxes["mean"][3]), ((boxes["white"][0] + boxes["white"][2]) // 2, boxes["white"][1]))
    for left, right in [("white", "spectro"), ("spectro", "regression"), ("regression", "output")]:
        _arrow(draw, ((boxes[left][2]), (boxes[left][1] + boxes[left][3]) // 2), (boxes[right][0], (boxes[right][1] + boxes[right][3]) // 2))

    _decision(
        draw,
        (1120, 825),
        (430, 250),
        "Validare",
        "daca RMSE_LOO < RMSE_baseline foloseste ridge",
    )
    _box(
        draw,
        (1400, 760, 1730, 935),
        "Fallback sigur",
        "daca modelul nu trece testul: S_hat = a*N + b sau media de antrenare",
        COLORS["red"],
        COLORS["red_edge"],
    )
    _arrow(draw, (1120, 620), (1120, 700))
    _arrow(draw, (1335, 825), (1400, 825), COLORS["red_edge"])
    _arrow(draw, (1290, 535), (1420, 760), COLORS["red_edge"])
    draw.text((955, 955), "Testul este leave-one-leaf-out: se scoate pe rand cate o frunza si se verifica eroarea predictiei.", font=small_font, fill=COLORS["muted"])

    formula_panel = (70, 730, 850, 1005)
    draw.rounded_rectangle(formula_panel, radius=20, fill="#ffffff", outline="#c7d0d9", width=3)
    draw.text((95, 755), "Formula finala in aplicatie", font=_font(28, bold=True), fill=COLORS["dark"])
    formulas = [
        "1) I_lambda(x,y) = max(R,G,B) / 255",
        "2) C_lambda = mean(I_lambda pe masca frunzei)",
        "3) N_lambda = C_lambda / W_lambda",
        "4) S_lambda = median(spectru in lambda +/- 10 nm)",
        "5) S_hat_lambda = b_lambda + SUM beta_lambda,k * ((N_k - mu_k) / sigma_k)",
        "6) foloseste modelul doar daca RMSE_LOO < RMSE_baseline",
    ]
    y = 805
    for formula in formulas:
        draw.text((100, y), formula, font=small_font, fill=COLORS["ink"])
        y += 31

    image.save(PNG_PATH, quality=95)
    return PNG_PATH


def _set_cell_shading(cell, fill: str) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill.replace("#", ""))
    cell._tc.get_or_add_tcPr().append(shading)


def _set_cell_text(cell, text: str, bold: bool = False, color: str = "1F2933") -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = Pt(9.5)
    run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def build_docx() -> Path:
    build_png()
    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    section.top_margin = Cm(1.1)
    section.bottom_margin = Cm(1.0)
    section.left_margin = Cm(1.15)
    section.right_margin = Cm(1.15)

    styles = doc.styles
    styles["Normal"].font.name = "Segoe UI"
    styles["Normal"].font.size = Pt(9.5)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run("Schema completa a calibrarii camera - spectrometru")
    title_run.bold = True
    title_run.font.size = Pt(18)
    title_run.font.color.rgb = RGBColor(24, 33, 47)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_run = subtitle.add_run("Fluxul de calcul folosit pentru pozele cu filtre si compararea cu spectrometrul")
    sub_run.font.size = Pt(10.5)
    sub_run.font.color.rgb = RGBColor(82, 97, 107)

    image_paragraph = doc.add_paragraph()
    image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    image_paragraph.add_run().add_picture(str(PNG_PATH), width=Inches(9.8))

    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = (Cm(4.0), Cm(8.0), Cm(14.0))
    headers = ("Etapa", "Formula", "Ce inseamna")
    for index, cell in enumerate(table.rows[0].cells):
        cell.width = widths[index]
        _set_cell_shading(cell, "#18212F")
        _set_cell_text(cell, headers[index], bold=True, color="FFFFFF")

    rows = [
        ("Pixel camera", "I_lambda(x,y) = max(R,G,B) / 255", "Poza filtrata este transformata intr-o intensitate pe pixel."),
        ("Medie frunza", "C_lambda = mean(I_lambda pe masca)", "Se folosesc doar pixelii detectati ca frunza."),
        ("Corectie alb", "N_lambda = C_lambda / W_lambda", "Compenseaza puterea diferita a filtrelor si a luminii din cutie."),
        ("Tinta spectrometru", "S_lambda = median(lambda +/- 10 nm)", "Reduce zgomotul fata de citirea unui singur punct spectral."),
        ("Regresie", "S_hat_lambda = b + SUM beta_k*((N_k-mu_k)/sigma_k)", "Foloseste toate benzile camerei cand exista suficiente filtre."),
        ("Validare", "RMSE_LOO < RMSE_baseline", "Modelul este folosit doar daca prezice mai bine decat media simpla."),
    ]
    for step, formula, meaning in rows:
        cells = table.add_row().cells
        for index, value in enumerate((step, formula, meaning)):
            cells[index].width = widths[index]
            _set_cell_text(cells[index], value, bold=index == 0)
            if len(table.rows) % 2 == 0:
                _set_cell_shading(cells[index], "#F3F6F8")

    note = doc.add_paragraph()
    note.paragraph_format.space_before = Pt(5)
    note.paragraph_format.space_after = Pt(0)
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    note_run = note.add_run(
        "Nota: referinta neagra este verificata, dar nu este scazuta cand ar produce valori fizic imposibile."
    )
    note_run.italic = True
    note_run.font.size = Pt(8.8)
    note_run.font.color.rgb = RGBColor(82, 97, 107)

    doc.save(DOCX_PATH)
    return DOCX_PATH


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    build_docx()
    print(DOCX_PATH)
    print(PNG_PATH)


if __name__ == "__main__":
    main()
