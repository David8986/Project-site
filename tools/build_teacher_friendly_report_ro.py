"""Build a short Romanian teacher-friendly report with many graphs."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"C:\Users\david\OneDrive\Desktop\Archive\Projects\CodeX")
REPORT_DIR = ROOT / "outputs" / "camera_validation_report"
GRAPH_DIR = REPORT_DIR / "graphs"
DOCX_PATH = REPORT_DIR / "raport_scurt_pentru_profesor_camera_spectrometru.docx"


FIGURES = [
    (
        "01_full_spectrum_raw_all.png",
        "Grafic 1. Toate spectrele masurate",
        "Aici se vad masuratorile brute de la spectrometru. Curbele nu sunt plate sau intamplatoare, deci probele chiar au semnal spectral diferit.",
    ),
    (
        "03_condition_mean_full_spectrum.png",
        "Grafic 2. Media spectrelor pentru fiecare tip de frunza",
        "Curbele medii arata ca frunzele sanatoase si cele nesanatoase nu raspund identic. Diferentele apar pe aproape tot spectrul, nu doar intr-un singur punct.",
    ),
    (
        "04_healthy_minus_unhealthy_difference.png",
        "Grafic 3. Diferenta dintre frunza sanatoasa si nesanatoasa",
        "Zonele verzi inseamna unde frunza sanatoasa are semnal mai mare. Zonele rosii inseamna unde frunza nesanatoasa are semnal mai mare.",
    ),
    (
        "17_spectrum_area_by_condition_filter.png",
        "Grafic 4. Raspunsul total al spectrometrului pe filtre",
        "Acest grafic rezuma tot spectrul pentru fiecare filtru. Se vede ca frunzele sanatoase si nesanatoase au forme de raspuns diferite.",
    ),
    (
        "07_camera_vs_spectrometer_normalized.png",
        "Grafic 5. Camera comparata cu spectrometrul",
        "Aici camera este comparata cu spectrometrul dupa normalizare. Ideea este sa vedem daca atunci cand spectrometrul creste, camera are un raspuns asemanator.",
    ),
    (
        "09_per_sample_correlation_bars.png",
        "Grafic 6. Cat de bine se potriveste camera cu spectrometrul",
        "Barele mai mari inseamna potrivire mai buna. Unele seturi se potrivesc bine, iar altele mai slab, probabil din cauza expunerii sau a pozitionarii diferite.",
    ),
    (
        "10_fingerprint_overlays_by_sample.png",
        "Grafic 7. Amprenta camerei si amprenta spectrometrului",
        "Pentru fiecare frunza, se compara forma raspunsului camerei cu forma raspunsului spectrometrului. Cand liniile au forme apropiate, camera confirma bine masuratoarea.",
    ),
    (
        "05_full_spectrum_pca.png",
        "Grafic 8. Gruparea spectrelor prin PCA",
        "PCA foloseste tot spectrul si il pune intr-un grafic 2D. Daca punctele se grupeaza, inseamna ca spectrele contin informatie reala despre probe.",
    ),
    (
        "11_camera_feature_heatmap.png",
        "Grafic 9. Harta valorilor masurate de camera",
        "Culorile arata cum se schimba valorile camerei intre filtre si probe. Nu toate imaginile sunt la fel, deci camera vede diferente intre masuratori.",
    ),
    (
        "12_spectrometer_feature_heatmap.png",
        "Grafic 10. Harta valorilor de la spectrometru",
        "Aceasta este varianta pentru spectrometru. Se poate compara vizual cu harta camerei ca sa vedem unde apar modele asemanatoare.",
    ),
    (
        "18_full_spectrum_similarity_heatmap.png",
        "Grafic 11. Cat de asemanatoare sunt spectrele intre ele",
        "Patratele mai deschise inseamna spectre mai asemanatoare. Acest grafic ajuta sa vedem daca masuratorile se grupeaza natural.",
    ),
    (
        "19_camera_contact_sheet.png",
        "Grafic 12. Imaginile folosite de camera",
        "Acesta este un tabel vizual cu pozele folosite. Se vede ca imaginile au fost facute cu filtre diferite, deci camera a primit informatie diferita pentru fiecare masuratoare.",
    ),
]


def main() -> None:
    doc = Document()
    setup(doc)
    add_title(doc)
    add_short_conclusion(doc)
    for index, (filename, title, explanation) in enumerate(FIGURES):
        add_figure(doc, GRAPH_DIR / filename, title, explanation, page_break_before=index > 0)
    add_final_note(doc)
    doc.save(DOCX_PATH)
    print(DOCX_PATH)


def setup(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)


def add_title(doc: Document) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Analiza camerei comparata cu spectrometrul")
    r.bold = True
    r.font.size = Pt(22)
    r.font.color.rgb = RGBColor(31, 78, 121)
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run("Raport scurt cu grafice si explicatii simple")
    r2.font.size = Pt(13)
    r2.font.color.rgb = RGBColor(90, 90, 90)


def add_short_conclusion(doc: Document) -> None:
    doc.add_paragraph()
    add_heading(doc, "Ideea pe scurt")
    points = [
        "Spectrometrul arata ca frunzele au raspunsuri diferite pe spectru.",
        "Camera nu da doar poze aleatorii: valorile ei se schimba in functie de filtru si de frunza.",
        "Pentru unele seturi, camera urmareste destul de bine raspunsul spectrometrului.",
        "Nu este inca o calibrare perfecta, dar este un proof of concept bun: camera poate produce date utile.",
    ]
    for point in points:
        doc.add_paragraph(point, style="List Bullet")


def add_figure(doc: Document, image_path: Path, title: str, explanation: str, page_break_before: bool = False) -> None:
    if page_break_before:
        doc.add_page_break()
    else:
        doc.add_paragraph()
    add_heading(doc, title)
    if image_path.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(image_path), width=Inches(6.7))
    else:
        doc.add_paragraph(f"Lipseste imaginea: {image_path.name}")
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p2.add_run(explanation)
    r.italic = True
    r.font.size = Pt(10.5)
    r.font.color.rgb = RGBColor(70, 70, 70)


def add_final_note(doc: Document) -> None:
    doc.add_paragraph()
    add_heading(doc, "Concluzie")
    doc.add_paragraph(
        "Concluzia mea este ca aceasta camera poate fi folosita pentru a obtine date reale despre frunze, "
        "mai ales daca este comparata cu spectrometrul pe aceleasi filtre. Pentru rezultate mai bune, ar trebui "
        "blocate expunerea si balansul de alb, iar fiecare masuratoare ar trebui repetata de mai multe ori."
    )


def add_heading(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(15)
    r.font.color.rgb = RGBColor(47, 98, 62)


if __name__ == "__main__":
    main()
