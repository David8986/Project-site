from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt
from lxml import etree


ROOT = Path(__file__).resolve().parents[1]
DOCX_PATH = ROOT / "fisa_formule_curent_continuu_BAC_RO.docx"
MML2OMML_PATH = Path(r"C:\Program Files\Microsoft Office\root\Office16\MML2OMML.XSL")
MATH_VENDOR = ROOT / "tools" / ".vendor_math"

if str(MATH_VENDOR) not in sys.path:
    sys.path.insert(0, str(MATH_VENDOR))

from latex2mathml.converter import convert as latex_to_mathml

PAGE_WIDTH = Mm(210)
PAGE_HEIGHT = Mm(297)
MARGIN_X = Mm(9)
MARGIN_Y = Mm(8)
OUTER_GAP = Mm(3)
BOX_PADDING = Mm(1.5)

FONT_TEXT = "Calibri"
FONT_MATH = "Cambria Math"
TITLE_SIZE = Pt(15)
SUBTITLE_SIZE = Pt(8.8)
SECTION_TITLE_SIZE = Pt(9.8)
BODY_SIZE = Pt(8.5)
SMALL_SIZE = Pt(8.1)


EQUATIONS = {
    "EQ_OHM_1": r"U = R I",
    "EQ_OHM_2": r"I = \frac{U}{R}",
    "EQ_OHM_3": r"R = \frac{U}{I}",
    "EQ_POWER_1": r"P = U I",
    "EQ_POWER_2": r"P = I^2 R",
    "EQ_POWER_3": r"P = \frac{U^2}{R}",
    "EQ_POWER_4": r"W = P t",
    "EQ_POWER_5": r"W = U I t",
    "EQ_POWER_6": r"W = I^2 R t",
    "EQ_POWER_7": r"W = \frac{U^2 t}{R}",
    "EQ_SERIES_1": r"I = I_1 = I_2 = \ldots",
    "EQ_SERIES_2": r"R_{eq} = R_1 + R_2 + \ldots",
    "EQ_SERIES_3": r"U = U_1 + U_2 + \ldots",
    "EQ_SERIES_4": r"U_k = I R_k",
    "EQ_PAR_1": r"U = U_1 = U_2 = \ldots",
    "EQ_PAR_2": r"I = I_1 + I_2 + \ldots",
    "EQ_PAR_3": r"\frac{1}{R_{eq}} = \frac{1}{R_1} + \frac{1}{R_2} + \ldots",
    "EQ_PAR_4": r"R_{eq} = \frac{R_1 R_2}{R_1 + R_2}",
    "EQ_PAR_5": r"I_k = \frac{U}{R_k}",
    "EQ_PAR_6": r"\frac{I_1}{I_2} = \frac{R_2}{R_1}",
    "EQ_KIR_1": r"\sum I_{in} = \sum I_{out}",
    "EQ_KIR_2": r"\sum U = 0",
    "EQ_GEN_1": r"U = E - I r",
    "EQ_GEN_2": r"U = E + I r",
    "EQ_GEN_3": r"I = \frac{E}{R + r}",
    "EQ_GEN_4": r"I_{sc} = \frac{E}{r}",
    "EQ_GEN_5": r"I = 0,\quad U = E",
    "EQ_GEN_6": r"R_{ext} = 0",
    "EQ_GS_1": r"E_{eq} = E_1 + E_2 + \ldots",
    "EQ_GS_2": r"r_{eq} = r_1 + r_2 + \ldots",
    "EQ_GS_3": r"E_{eq} = E_1 - E_2",
    "EQ_GS_4": r"r_{eq} = r_1 + r_2",
    "EQ_GP_1": r"U_g = \frac{\frac{E_1}{r_1} + \frac{E_2}{r_2}}{\frac{1}{r_1} + \frac{1}{r_2}}",
    "EQ_GP_2": r"r_{eq} = \frac{r_1 r_2}{r_1 + r_2}",
    "EQ_GP_3": r"U_g = \frac{\frac{E_1}{r_1} - \frac{E_2}{r_2}}{\frac{1}{r_1} + \frac{1}{r_2}}",
    "EQ_REO_1": r"R_{BD} = R",
    "EQ_REO_2": r"R_{BC} + R_{CD} = R",
}

MATH_TRANSFORM = etree.XSLT(etree.parse(str(MML2OMML_PATH)))


def set_cell_margins(cell, top=None, start=None, bottom=None, end=None):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        if value is None:
            continue
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.first_child_found_in("w:shd")
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)


def set_cell_borders(cell, color="7F7F7F", size=6):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge in ("top", "left", "bottom", "right"):
        elem = tc_borders.find(qn(f"w:{edge}"))
        if elem is None:
            elem = OxmlElement(f"w:{edge}")
            tc_borders.append(elem)
        elem.set(qn("w:val"), "single")
        elem.set(qn("w:sz"), str(size))
        elem.set(qn("w:space"), "0")
        elem.set(qn("w:color"), color)


def set_table_borders_none(table):
    tbl_pr = table._tbl.tblPr
    tbl_borders = tbl_pr.first_child_found_in("w:tblBorders")
    if tbl_borders is None:
        tbl_borders = OxmlElement("w:tblBorders")
        tbl_pr.append(tbl_borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        elem = tbl_borders.find(qn(f"w:{edge}"))
        if elem is None:
            elem = OxmlElement(f"w:{edge}")
            tbl_borders.append(elem)
        elem.set(qn("w:val"), "nil")


def emu_to_twips(value):
    return int(round(int(value) / 635))


def set_font(run, name=FONT_TEXT, size=BODY_SIZE, bold=False, italic=False):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run.font.size = size
    run.bold = bold
    run.italic = italic


def omml_for_latex(latex):
    mathml = latex_to_mathml(latex)
    math_tree = etree.fromstring(mathml.encode("utf-8"))
    omml_tree = MATH_TRANSFORM(math_tree)
    return omml_tree.getroot()


def style_paragraph(paragraph, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=0.6):
    paragraph.alignment = align
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(space_after)
    fmt.line_spacing = 1.0
    fmt.left_indent = Mm(0)
    fmt.right_indent = Mm(0)


class CellWriter:
    def __init__(self, cell):
        self.cell = cell
        self.first_used = False

    def _paragraph(self):
        if not self.first_used:
            self.first_used = True
            return self.cell.paragraphs[0]
        return self.cell.add_paragraph()

    def text(self, text, *, bold=False, italic=False, size=BODY_SIZE, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=0.6):
        p = self._paragraph()
        style_paragraph(p, align=align, space_after=space_after)
        run = p.add_run(text)
        set_font(run, size=size, bold=bold, italic=italic)
        return p

    def equation(self, placeholder, *, size=BODY_SIZE, space_after=0.6):
        p = self._paragraph()
        style_paragraph(p, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=space_after)
        latex = EQUATIONS[placeholder]
        run = p.add_run()
        set_font(run, name=FONT_MATH, size=size)
        p._p.append(omml_for_latex(latex))
        return p


def configure_page(section):
    section.page_width = PAGE_WIDTH
    section.page_height = PAGE_HEIGHT
    section.left_margin = MARGIN_X
    section.right_margin = MARGIN_X
    section.top_margin = MARGIN_Y
    section.bottom_margin = MARGIN_Y
    section.header_distance = Mm(5)
    section.footer_distance = Mm(5)


def apply_base_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = FONT_TEXT
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT_TEXT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_TEXT)
    normal.font.size = BODY_SIZE


def set_table_width(table, width_twips):
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(width_twips))
    tbl_w.set(qn("w:type"), "dxa")


def add_title(doc):
    p = doc.add_paragraph()
    style_paragraph(p, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=0)
    run = p.add_run("Fișă de formule BAC - Curent continuu / Circuite electrice")
    set_font(run, size=TITLE_SIZE, bold=True)

    p = doc.add_paragraph()
    style_paragraph(p, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
    run = p.add_run("Formule esențiale, observații de semn și pași rapizi pentru probleme de tip Bacalaureat")
    set_font(run, size=SUBTITLE_SIZE, italic=True)


def add_page_kicker(doc, text):
    p = doc.add_paragraph()
    style_paragraph(p, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=3)
    run = p.add_run(text)
    set_font(run, size=SUBTITLE_SIZE, bold=True)


def build_page_table(doc):
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders_none(table)

    section = doc.sections[-1]
    usable_width = section.page_width - section.left_margin - section.right_margin
    set_table_width(table, emu_to_twips(usable_width))

    left, right = table.rows[0].cells
    column_width = int((usable_width - int(OUTER_GAP)) / 2)
    left.width = column_width
    right.width = column_width
    for cell in (left, right):
        cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
        set_cell_margins(cell, top=0, bottom=0)
    set_cell_margins(left, end=int(OUTER_GAP.twips))
    set_cell_margins(right, start=int(OUTER_GAP.twips))
    return left, right


def add_section_box(container_cell, title):
    box = container_cell.add_table(rows=2, cols=1)
    box.alignment = WD_TABLE_ALIGNMENT.CENTER
    box.autofit = False
    container_width = container_cell.width or 0
    set_table_width(box, max(int(emu_to_twips(container_width) * 0.98), 5000))

    title_cell = box.rows[0].cells[0]
    body_cell = box.rows[1].cells[0]
    for cell in (title_cell, body_cell):
        cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
        set_cell_borders(cell)
    set_cell_shading(title_cell, "D9D9D9")
    set_cell_margins(title_cell, top=40, start=90, bottom=40, end=90)
    set_cell_margins(body_cell, top=90, start=110, bottom=90, end=110)

    p = title_cell.paragraphs[0]
    style_paragraph(p, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=0)
    run = p.add_run(title)
    set_font(run, size=SECTION_TITLE_SIZE, bold=True)
    return CellWriter(body_cell)


def add_spacer(cell, height_pt=3):
    p = cell.add_paragraph()
    style_paragraph(p, space_after=0)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(height_pt)


def add_units_section(cell):
    writer = add_section_box(cell, "1. Mărimi și unități")
    writer.text("Simboluri de bază folosite în problemele de circuit:", bold=True, size=SMALL_SIZE, space_after=1.2)

    table = writer.cell.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    table.autofit = False
    headers = ("Simbol", "Mărime", "Unitate")
    widths = (Mm(16), Mm(47), Mm(18))
    for idx, header in enumerate(headers):
        hdr = table.rows[0].cells[idx]
        hdr.width = widths[idx]
        set_cell_shading(hdr, "F2F2F2")
        set_cell_margins(hdr, top=45, start=70, bottom=45, end=70)
        hdr.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = hdr.paragraphs[0]
        style_paragraph(p, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=0)
        run = p.add_run(header)
        set_font(run, size=SMALL_SIZE, bold=True)

    rows = [
        ("I", "intensitatea curentului electric", "A"),
        ("U", "tensiunea electrică", "V"),
        ("R", "rezistența electrică", "Ω"),
        ("P", "puterea electrică", "W"),
        ("W / E_el", "energia electrică", "J"),
        ("E", "t.e.m. a generatorului", "V"),
        ("r", "rezistența internă", "Ω"),
    ]
    for row_values in rows:
        row = table.add_row()
        for idx, value in enumerate(row_values):
            cell = row.cells[idx]
            cell.width = widths[idx]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell, top=35, start=70, bottom=35, end=70)
            p = cell.paragraphs[0]
            align = WD_ALIGN_PARAGRAPH.CENTER if idx != 1 else WD_ALIGN_PARAGRAPH.LEFT
            style_paragraph(p, align=align, space_after=0)
            run = p.add_run(value)
            set_font(run, size=SMALL_SIZE)


def add_ohm_section(cell):
    writer = add_section_box(cell, "2. Legea lui Ohm")
    writer.equation("EQ_OHM_1")
    writer.equation("EQ_OHM_2")
    writer.equation("EQ_OHM_3")
    writer.text(
        "Pentru un conductor sau consumator, tensiunea scade în sensul curentului convențional.",
        size=SMALL_SIZE,
        space_after=0,
    )


def add_power_section(cell):
    writer = add_section_box(cell, "3. Putere și energie electrică")
    for key in (
        "EQ_POWER_1",
        "EQ_POWER_2",
        "EQ_POWER_3",
        "EQ_POWER_4",
        "EQ_POWER_5",
        "EQ_POWER_6",
        "EQ_POWER_7",
    ):
        writer.equation(key)


def add_resistors_section(cell):
    writer = add_section_box(cell, "4. Gruparea rezistoarelor")
    writer.text("Serie", bold=True, size=SMALL_SIZE, space_after=0.4)
    for key in ("EQ_SERIES_1", "EQ_SERIES_2", "EQ_SERIES_3", "EQ_SERIES_4"):
        writer.equation(key)
    writer.text("Paralel", bold=True, size=SMALL_SIZE, space_after=0.4)
    for key in ("EQ_PAR_1", "EQ_PAR_2", "EQ_PAR_3", "EQ_PAR_4", "EQ_PAR_5", "EQ_PAR_6"):
        size = SMALL_SIZE if key in {"EQ_PAR_3", "EQ_PAR_4", "EQ_PAR_6"} else BODY_SIZE
        writer.equation(key, size=size)


def add_kirchhoff_section(cell):
    writer = add_section_box(cell, "5. Legile lui Kirchhoff")
    writer.text("Legea nodurilor", bold=True, size=SMALL_SIZE, space_after=0.4)
    writer.equation("EQ_KIR_1")
    writer.text("Curentul se împarte sau se reunește numai în noduri.", size=SMALL_SIZE, space_after=1.1)
    writer.text("Legea ochiurilor", bold=True, size=SMALL_SIZE, space_after=0.4)
    writer.equation("EQ_KIR_2")
    writer.text("Suma creșterilor de tensiune este egală cu suma căderilor de tensiune.", size=SMALL_SIZE)
    writer.text("Pe consumatori apare cădere de tensiune, iar pe surse de la minus la plus apare creștere.", size=SMALL_SIZE, space_after=0)


def add_meters_section(cell):
    writer = add_section_box(cell, "8. Aparate ideale de măsură")
    writer.text("Ampermetru", bold=True, size=SMALL_SIZE, space_after=0.2)
    writer.text("- R_A ≈ 0", size=SMALL_SIZE, space_after=0.1)
    writer.text("- se leagă în serie", size=SMALL_SIZE, space_after=0.1)
    writer.text("- se comportă ca un fir ideal", size=SMALL_SIZE, space_after=0.8)
    writer.text("Voltmetru", bold=True, size=SMALL_SIZE, space_after=0.2)
    writer.text("- R_V → ∞", size=SMALL_SIZE, space_after=0.1)
    writer.text("- se leagă în paralel", size=SMALL_SIZE, space_after=0.1)
    writer.text("- printr-un voltmetru ideal nu trece curent", size=SMALL_SIZE, space_after=0)


def add_reostat_section(cell):
    writer = add_section_box(cell, "9. Reostat / potențiometru")
    writer.equation("EQ_REO_1")
    writer.equation("EQ_REO_2")
    writer.text("Cursorul C împarte rezistorul în două porțiuni care pot funcționa separat în circuit.", size=SMALL_SIZE)
    writer.text("Punctele unite prin fir ideal au același potențial.", size=SMALL_SIZE, space_after=0)


def add_generator_section(cell):
    writer = add_section_box(cell, "6. Generator real")
    writer.equation("EQ_GEN_1")
    writer.text("Când generatorul debitează curent.", size=SMALL_SIZE, space_after=0.8)
    writer.equation("EQ_GEN_2")
    writer.text("Când generatorul este încărcat și primește curent.", size=SMALL_SIZE, space_after=0.8)
    writer.equation("EQ_GEN_3")
    writer.equation("EQ_GEN_4")
    writer.equation("EQ_GEN_5")
    writer.text("Circuit deschis", bold=True, size=SMALL_SIZE, space_after=0.2)
    writer.text("I = 0 și tensiunea la borne este egală cu t.e.m.", size=SMALL_SIZE, space_after=0.8)
    writer.text("Scurtcircuit", bold=True, size=SMALL_SIZE, space_after=0.2)
    writer.equation("EQ_GEN_6", size=SMALL_SIZE)


def add_generator_group_section(cell):
    writer = add_section_box(cell, "7. Gruparea generatoarelor")
    writer.text("Serie, același sens", bold=True, size=SMALL_SIZE, space_after=0.3)
    writer.equation("EQ_GS_1")
    writer.equation("EQ_GS_2")
    writer.text("Serie, sensuri opuse", bold=True, size=SMALL_SIZE, space_after=0.3)
    writer.equation("EQ_GS_3")
    writer.equation("EQ_GS_4")
    writer.text("Paralel, aceeași polaritate (două generatoare)", bold=True, size=SMALL_SIZE, space_after=0.3)
    writer.equation("EQ_GP_1", size=SMALL_SIZE)
    writer.equation("EQ_GP_2")
    writer.text("Paralel, polarități opuse", bold=True, size=SMALL_SIZE, space_after=0.3)
    writer.equation("EQ_GP_3", size=SMALL_SIZE)
    writer.text("Semnul lui U_g arată polaritatea reală a grupării.", size=SMALL_SIZE, space_after=0)


def add_rules_section(cell):
    writer = add_section_box(cell, "10. Reguli utile la probleme de circuit")
    rules = [
        "Curentul convențional are sensul de la plus la minus în circuitul exterior.",
        "La baterie, bara lungă este plus, iar bara scurtă este minus.",
        "Pe un fir ideal, tensiunea nu se modifică.",
        "Într-un nod, curentul se poate împărți sau reuni.",
        "Prin orice componentă dintr-o ramură, curentul care intră este egal cu cel care iese.",
        "Pe un consumator, tensiunea scade în sensul curentului.",
        "În serie: același curent. În paralel: aceeași tensiune.",
        "Dacă există o ramură de scurtcircuit, curentul preferă ramura cu R = 0.",
        "Dacă o ramură este întreruptă, curentul prin acea ramură este 0.",
    ]
    for item in rules:
        writer.text(f"- {item}", size=SMALL_SIZE, space_after=0.2)


def add_checklist_section(cell):
    writer = add_section_box(cell, "Checklist - cum rezolvi rapid")
    steps = [
        "Marchează plusul și minusul fiecărei surse.",
        "Alege sensurile curenților convenționali pe ramuri.",
        "Înlocuiește ampermetrele ideale cu fire și voltmetrele ideale cu ramuri deschise.",
        "Identifică nodurile, ramurile și eventualele scurtcircuite.",
        "Reduce grupările serie și paralel acolo unde se poate.",
        "Aplică legile lui Kirchhoff dacă schema nu se reduce direct.",
        "Verifică la final unitățile și semnele tensiunilor și curenților.",
    ]
    for index, step in enumerate(steps, start=1):
        writer.text(f"{index}. {step}", size=SMALL_SIZE, space_after=0.35)


def add_exam_notes_section(cell):
    writer = add_section_box(cell, "Observații rapide de examen")
    notes = [
        "Nu confunda t.e.m. E cu tensiunea la borne U a generatorului real.",
        "Dacă un rezistor este ocolit de un fir ideal, tensiunea pe el devine 0.",
        "În paralel, rezistența echivalentă este mai mică decât cea mai mică rezistență din grup.",
        "I_sc este curentul maxim al generatorului real și apare pentru R_ext = 0.",
        "Când semnul rezultatului nu corespunde ipotezei tale de sens, curentul real are sens opus.",
    ]
    for item in notes:
        writer.text(f"- {item}", size=SMALL_SIZE, space_after=0.25)


def build_document():
    doc = Document()
    configure_page(doc.sections[0])
    apply_base_styles(doc)

    add_title(doc)

    left, right = build_page_table(doc)
    add_units_section(left)
    add_spacer(left)
    add_ohm_section(left)
    add_spacer(left)
    add_power_section(left)
    add_spacer(left)
    add_kirchhoff_section(left)

    add_resistors_section(right)
    add_spacer(right)
    add_meters_section(right)
    add_spacer(right)
    add_reostat_section(right)

    doc.add_page_break()
    add_page_kicker(doc, "Generator real, grupări de surse și strategie de rezolvare")

    left, right = build_page_table(doc)
    add_generator_section(left)
    add_spacer(left)
    add_generator_group_section(left)

    add_rules_section(right)
    add_spacer(right)
    add_checklist_section(right)
    add_spacer(right)
    add_exam_notes_section(right)

    doc.save(DOCX_PATH)


if __name__ == "__main__":
    build_document()
