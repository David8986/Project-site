from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph


DOCX_PATH = Path(r"C:\Users\david\OneDrive\Desktop\referat modificat FINAL.docx")


def has_page_break(paragraph_element) -> bool:
    if paragraph_element is None:
        return False
    for br in paragraph_element.iter():
        if br.tag.endswith("}br") and br.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}type") == "page":
            return True
    return False


def insert_page_break_before(paragraph: Paragraph) -> None:
    previous = paragraph._p.getprevious()
    if has_page_break(previous):
        return
    new_p = OxmlElement("w:p")
    paragraph._p.addprevious(new_p)
    break_paragraph = Paragraph(new_p, paragraph._parent)
    break_paragraph.add_run().add_break(WD_BREAK.PAGE)


def main() -> None:
    document = Document(DOCX_PATH)
    for paragraph in document.paragraphs:
        if paragraph.text.strip().startswith("Figura 5. Harta de similaritate"):
            insert_page_break_before(paragraph)
            document.save(DOCX_PATH)
            print(f"Updated layout in: {DOCX_PATH}")
            return
    raise RuntimeError("Nu am gasit legenda pentru Figura 5.")


if __name__ == "__main__":
    main()
