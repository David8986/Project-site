from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE_DOCX = Path(r"D:\downloads\referat (2) - cu date noi.docx")
OUTPUT_DOCX = Path(r"D:\downloads\referat (2) - revizuit coerent.docx")
SUMMARY_PATH = ROOT / "outputs" / "inside_outside_comparison" / "inside_outside_summary.json"
GRAPH_DIR = ROOT / "outputs" / "inside_outside_comparison" / "graphs"


FIGURES = [
    {
        "file": "01_mean_normalized_spectra_inside_vs_outside.png",
        "title": "Figura 1. Compararea formei spectrelor normalizate",
        "width": 5.85,
        "text": (
            "În această reprezentare, spectrele măsurate afară și în cutie sunt normalizate, "
            "astfel încât comparația să nu fie dominată de diferența de intensitate a luminii. "
            "Se observă că mediul de măsurare modifică forma exactă a curbei, mai ales în zona "
            "vizibilă, dar păstrează aceeași tendință generală: scăderea și creșterea răspunsului "
            "apar în regiuni apropiate ale spectrului. Acest lucru susține ideea că sistemul "
            "camerei poate urmări variații reale ale răspunsului spectral, cu condiția folosirii "
            "unei calibrări corecte."
        ),
    },
    {
        "file": "03_similarity_by_filter_and_condition.png",
        "title": "Figura 2. Similaritatea pe filtre și pe starea frunzei",
        "width": 5.65,
        "text": (
            "Graficul arată cât de apropiată este forma spectrului măsurat în cutie față de "
            "spectrul măsurat afară, pentru fiecare filtru și pentru cele două categorii de frunze. "
            "Coeficientul Pearson r este mai mare atunci când forma curbelor se potrivește mai bine. "
            "Valorile cele mai utile pentru validarea prototipului apar în special în zona red-edge "
            "și NIR, unde unele filtre păstrează o asemănare moderată sau bună. Valorile mai mici "
            "arată că anumite lungimi de undă sunt mai sensibile la iluminare și la geometria montajului."
        ),
    },
    {
        "file": "04_raw_signal_area_inside_vs_outside.png",
        "title": "Figura 3. Diferența dintre intensitatea brută măsurată afară și în cutie",
        "width": 5.65,
        "page_break_before": True,
        "text": (
            "Acest grafic compară aria totală a semnalului spectral, adică energia brută înregistrată "
            "pe întregul spectru. Semnalul obținut afară este de ordinul sutelor de ori mai mare decât "
            "semnalul obținut în cutie, mediana raportului fiind aproximativ 137,8. Din acest motiv, "
            "valorile brute nu pot fi comparate direct ca și cum ar proveni din același mediu de lucru. "
            "Pentru interpretare științifică trebuie folosite normalizarea, raporturile între benzi și "
            "comparația formei curbelor."
        ),
    },
    {
        "file": "05_noise_roughness_inside_vs_outside.png",
        "title": "Figura 4. Zgomotul și rugozitatea curbelor spectrale",
        "width": 5.5,
        "text": (
            "Figura compară stabilitatea curbelor prin doi indicatori: zgomotul relativ al semnalului "
            "și rugozitatea, adică variațiile rapide de la un punct spectral la următorul. În cutie, "
            "semnalul este mai slab, iar variațiile mici devin mai vizibile după prelucrare. Totuși, "
            "acest lucru nu invalidează metoda, ci arată că sistemul are nevoie de expunere controlată, "
            "calibrare și repetarea măsurătorilor. Pentru camera multispectrală, aceste rezultate indică "
            "unde trebuie îmbunătățită iluminarea internă."
        ),
    },
    {
        "file": "02_inside_outside_similarity_heatmap.png",
        "title": "Figura 5. Harta de similaritate pentru perechile spectrale",
        "width": 3.55,
        "page_break_before": True,
        "text": (
            "Harta de similaritate prezintă fiecare pereche de măsurători afară-în cutie. Culorile verzi "
            "indică perechi cu o asemănare mai bună a formei spectrului, iar culorile galben-roșii arată "
            "perechi în care mediul de măsurare a schimbat mai puternic răspunsul. Mediana generală de "
            "aproximativ r = 0,56 nu înseamnă că sistemul este identic cu spectrometrul în orice condiție, "
            "ci că există o relație măsurabilă și utilă între cele două tipuri de achiziție. Această relație "
            "poate fi folosită ca bază pentru calibrarea aplicației."
        ),
    },
]


def body_text(element) -> str:
    return "".join(t.text or "" for t in element.iter(qn("w:t"))).strip()


def remove_previous_added_section(doc: Document) -> None:
    body = doc.element.body
    children = list(body)
    start = None
    markers = [
        "Anexa experimentala: comparatia masuratorilor afara / in cutie",
        "Anexa experimentală: comparația măsurătorilor afară / în cutie",
    ]
    for idx, child in enumerate(children):
        text = body_text(child)
        if any(marker in text for marker in markers):
            start = idx
            break
    if start is None:
        raise RuntimeError("Could not find the previously inserted experimental section.")

    while start > 0 and not body_text(children[start - 1]):
        start -= 1

    end = start
    while end < len(children) and children[end].tag != qn("w:sectPr"):
        end += 1

    for child in children[start:end]:
        body.remove(child)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text: str, bold: bool = False) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(9)
    run.bold = bold
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "BFCABC")


def add_compact_paragraph(doc: Document, text: str, size: float = 10.5, bold_prefix: str | None = None) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.05
    if bold_prefix and text.startswith(bold_prefix):
        run = paragraph.add_run(bold_prefix)
        run.bold = True
        run.font.name = "Arial"
        run.font.size = Pt(size)
        rest = text[len(bold_prefix) :]
        run = paragraph.add_run(rest)
    else:
        run = paragraph.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(size)


def add_heading(doc: Document, text: str, level: int) -> None:
    paragraph = doc.add_paragraph()
    paragraph.style = f"Heading {level}"
    paragraph.paragraph_format.space_before = Pt(6 if level == 1 else 4)
    paragraph.paragraph_format.space_after = Pt(5)
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run(text)
    run.font.name = "Arial"
    run.font.color.rgb = RGBColor(23, 33, 27)
    run.bold = True
    run.font.size = Pt(15 if level == 1 else 12)


def add_summary_table(doc: Document, summary: dict) -> None:
    rows = [
        ("Perechi spectrale comparate", str(summary["matched_inside_outside_spectrum_pairs"]), "măsurători afară corelate cu măsurători în cutie"),
        ("Similaritate mediană totală", "r = 0,56", "asemănare moderată a formei spectrale"),
        ("Similaritate frunze sănătoase", "r = 0,57", "rezultat apropiat de media generală"),
        ("Similaritate frunze nesănătoase", "r = 0,53", "asemănare moderată, cu variații mai mari"),
        ("Raport median al semnalului", "137,8x", "semnalul brut este mult mai puternic afară"),
    ]
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table)
    header = table.rows[0].cells
    for cell, text in zip(header, ["Indicator", "Valoare", "Interpretare"]):
        set_cell_text(cell, text, bold=True)
        set_cell_shading(cell, "DDEBDD")
    for indicator, value, interpretation in rows:
        cells = table.add_row().cells
        set_cell_text(cells[0], indicator)
        set_cell_text(cells[1], value)
        set_cell_text(cells[2], interpretation)
    for row in table.rows:
        row.cells[0].width = Inches(2.1)
        row.cells[1].width = Inches(1.2)
        row.cells[2].width = Inches(3.0)
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(3)


def add_figure(doc: Document, figure: dict) -> None:
    title = doc.add_paragraph()
    title.paragraph_format.page_break_before = bool(figure.get("page_break_before"))
    title.paragraph_format.space_before = Pt(5)
    title.paragraph_format.space_after = Pt(2)
    title.paragraph_format.keep_with_next = True
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run(figure["title"])
    title_run.bold = True
    title_run.font.name = "Arial"
    title_run.font.size = Pt(9.5)

    image_paragraph = doc.add_paragraph()
    image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    image_paragraph.paragraph_format.space_after = Pt(2)
    image_paragraph.paragraph_format.keep_with_next = True
    image_paragraph.add_run().add_picture(str(GRAPH_DIR / figure["file"]), width=Inches(figure["width"]))

    explanation = doc.add_paragraph()
    explanation.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    explanation.paragraph_format.space_after = Pt(4)
    explanation.paragraph_format.line_spacing = 1.0
    run = explanation.add_run(figure["text"])
    run.font.name = "Arial"
    run.font.size = Pt(8.7)


def add_rewritten_section(doc: Document) -> None:
    summary = json.loads(SUMMARY_PATH.read_text(encoding="utf-8"))

    add_heading(doc, "VI. Validarea experimentală a sistemului prin comparație spectrometrică", 1)
    add_compact_paragraph(
        doc,
        "Pentru verificarea practică a prototipului, măsurătorile spectrometrice realizate în exterior au fost comparate cu măsurătorile realizate în cutia de testare. Comparația a fost făcută pe perechi cu aceeași stare a frunzei și același filtru spectral, folosind întregul spectru măsurat, nu doar câteva puncte izolate.",
    )
    add_compact_paragraph(
        doc,
        "Scopul acestei etape nu este ca valorile brute să fie identice, deoarece iluminarea și geometria de măsurare sunt diferite. Important este dacă forma răspunsului spectral și tendințele dintre filtre se păstrează suficient de bine pentru ca sistemul optic cu cameră să poată fi calibrat și folosit ca soluție accesibilă de analiză multispectrală.",
    )
    add_compact_paragraph(
        doc,
        "Fotografiile camerei sunt folosite în această etapă pentru seria realizată în cutie, iar măsurările din exterior funcționează ca reper spectrometric. De aceea, comparația de mai jos urmărește mai ales schimbarea produsă de mediul de măsurare, pentru a stabili ce corecții trebuie aplicate înainte ca rezultatele camerei să fie comparate direct cu exteriorul.",
    )
    add_compact_paragraph(
        doc,
        "În tabel sunt sintetizate rezultatele principale ale comparației. Coeficientul Pearson r descrie asemănarea formei curbelor după normalizare, iar raportul de semnal arată diferența dintre intensitatea brută înregistrată afară și cea obținută în cutie.",
    )
    add_summary_table(doc, summary)

    for figure in FIGURES:
        add_figure(doc, figure)

    add_heading(doc, "Interpretarea rezultatelor experimentale", 2)
    add_compact_paragraph(
        doc,
        "Rezultatele arată că măsurătorile din cutie și cele din exterior nu sunt identice, însă păstrează o legătură măsurabilă. Similaritatea mediană de aproximativ r = 0,56 indică o asemănare moderată a formei spectrale, ceea ce este important pentru un prototip low-cost aflat în etapa de validare. Diferența cea mai mare este intensitatea brută a semnalului, deoarece lumina naturală din exterior este mult mai puternică decât iluminarea din cutie.",
    )
    add_compact_paragraph(
        doc,
        "Prin urmare, sistemul nu trebuie evaluat doar prin valori brute de luminozitate, ci prin comparații normalizate, rapoarte între benzi și comportamentul relativ al curbelor. Această concluzie este coerentă cu obiectivul proiectului: realizarea unei camere multispectrale accesibile, care să poată aproxima informații utile despre starea plantei și să fie îmbunătățită prin calibrare, iluminare controlată și repetarea măsurătorilor.",
    )


def main() -> None:
    doc = Document(SOURCE_DOCX)
    remove_previous_added_section(doc)
    add_rewritten_section(doc)
    doc.save(OUTPUT_DOCX)
    print(OUTPUT_DOCX)


if __name__ == "__main__":
    main()
