"""Build a concise DOCX summary of recent app changes."""

from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


PROJECT_ROOT = Path(__file__).resolve().parents[1]
PROFILE_PATH = PROJECT_ROOT / "plant_health_mvp_new_data" / "data" / "real_validation" / "camera_calibration_profile.json"
REPORT_PATH = (
    PROJECT_ROOT
    / "plant_health_mvp_new_data"
    / "runs"
    / "user_check_healthy_yellow_2026_05_24_fixed_mask"
    / "mapping_report.json"
)
OUTPUT = Path(r"C:\Users\david\OneDrive\Desktop\rezumat_schimbari_app_date_noi.docx")


ACCENT = "1F5F7A"
LIGHT = "EAF2F8"
MID = "D6E6F2"
TEXT = RGBColor(30, 30, 30)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def style_table(table, header_fill: str = ACCENT) -> None:
    table.style = "Table Grid"
    for row_index, row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.08
                for run in paragraph.runs:
                    run.font.name = "Aptos"
                    run.font.size = Pt(9)
            if row_index == 0:
                set_cell_shading(cell, header_fill)
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)
            else:
                set_cell_shading(cell, "FFFFFF" if row_index % 2 else "F7FAFC")


def set_widths(table, widths_cm: list[float]) -> None:
    for row in table.rows:
        for cell, width in zip(row.cells, widths_cm):
            cell.width = Cm(width)


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    heading = doc.add_heading(text, level=level)
    heading.paragraph_format.space_before = Pt(9 if level == 1 else 6)
    heading.paragraph_format.space_after = Pt(4)
    for run in heading.runs:
        run.font.name = "Aptos Display"
        run.font.color.rgb = RGBColor.from_string(ACCENT)
        run.font.bold = True


def add_callout(doc: Document, title: str, body: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.autofit = True
    cell = table.cell(0, 0)
    set_cell_shading(cell, LIGHT)
    set_cell_margins(cell, top=140, bottom=140, start=160, end=160)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(title)
    r.font.bold = True
    r.font.color.rgb = RGBColor.from_string(ACCENT)
    r.font.size = Pt(10.5)
    p2 = cell.add_paragraph(body)
    p2.paragraph_format.space_after = Pt(0)
    for run in p2.runs:
        run.font.size = Pt(9.5)


def bullet(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.left_indent = Cm(0.45)
    run = p.add_run(text)
    run.font.name = "Aptos"
    run.font.size = Pt(9.5)
    run.font.color.rgb = TEXT


def load_numbers() -> dict[str, object]:
    profile = json.loads(PROFILE_PATH.read_text(encoding="utf-8"))
    report = json.loads(REPORT_PATH.read_text(encoding="utf-8"))
    source_940 = profile["source_response_compensation"]["bands"]["940"]
    corrected_940 = report["camera_calibration"]["corrected_bands"]["940"]
    return {
        "source_940": source_940,
        "corrected_940": corrected_940,
        "profile_id": profile["profile_id"],
        "mask_method": report["vegetation"]["mask"].get("method"),
        "mask_used_bands": ", ".join(report["vegetation"]["mask"].get("used_bands", [])),
        "vegetation_pixels": report["vegetation"]["spectrum_metadata"].get("vegetation_pixel_count"),
        "pytest": "6 tests passed",
    }


def build_doc() -> Path:
    data = load_numbers()
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.62)
    section.right_margin = Inches(0.62)

    styles = doc.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(9.5)
    styles["Normal"].font.color.rgb = TEXT

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(2)
    run = title.add_run("Rezumat schimbari app - date reale")
    run.font.name = "Aptos Display"
    run.font.size = Pt(20)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string(ACCENT)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(8)
    r = subtitle.add_run("Ce s-a schimbat de cand am trecut de la demo la masuratorile reale camera + spectrometru")
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor(90, 90, 90)

    add_callout(
        doc,
        "Ideea principala",
        (
            "App-ul nu mai este doar o demonstratie. Acum foloseste imaginile reale pe filtre, "
            "referintele alb/negru, curba sursei de lumina si comparatia cu spectrometrul ca sa scoata "
            "valori calibrate si valori de incredere."
        ),
    )

    add_heading(doc, "1. Schimbari principale", 1)
    table = doc.add_table(rows=1, cols=3)
    table.rows[0].cells[0].text = "Zona"
    table.rows[0].cells[1].text = "Ce s-a schimbat"
    table.rows[0].cells[2].text = "De ce conteaza"
    rows = [
        (
            "Date reale",
            "App-ul citeste seturile reale din folderul data analisys si pastreaza output-uri separate pentru rularea noua.",
            "Analiza se bazeaza pe masuratori reale, nu pe date simulate.",
        ),
        (
            "Import imagini",
            "Wizard-ul poate deschide un JSON existent, modifica fisierele si inlocui o imagine fara sa refaci totul de la zero.",
            "Poti corecta inputul rapid cand o poza este gresita sau vrei alt filtru.",
        ),
        (
            "Masca frunzei",
            "Detectia frunzei foloseste acum foreground pe benzile filtrate, threshold robust, morfologie si cea mai plauzibila componenta.",
            "Nu mai selecteaza lumina verde de fundal; se apropie mult mai bine de forma frunzei.",
        ),
        (
            "Calibrare camera",
            "Valorile brute sunt normalizate cu referinta alba si apoi comparate cu ferestre spectrale ale spectrometrului.",
            "Camera devine comparabila numeric cu spectrometrul, nu doar vizual.",
        ),
        (
            "Modele spectrometru",
            "Pentru fiecare lungime de unda, app-ul incearca regresie ridge multibanda si foloseste fallback single-band cand validarea nu e buna.",
            "Evita modele care par bune doar pentru ca avem putine probe.",
        ),
        (
            "940 nm",
            "Banda 940 ramane vizibila, dar este marcata low-confidence pentru concluzii calibrate.",
            "Nu lasam o banda zgomotoasa sa strice concluzia proiectului.",
        ),
    ]
    for item in rows:
        cells = table.add_row().cells
        for cell, text in zip(cells, item):
            cell.text = text
    set_widths(table, [3.1, 7.0, 6.4])
    style_table(table)

    add_heading(doc, "2. Algoritmul actual", 1)
    bullets = [
        f"Masca folosita la proba curenta: {data['mask_method']} pe banda/banzile {data['mask_used_bands']}.",
        f"Numar pixeli frunza in rularea curenta: {data['vegetation_pixels']}.",
        "Pentru fiecare filtru se calculeaza media pixelilor frunzei, nu media intregii imagini.",
        "Camera se normalizeaza cu alb: camera_alb = camera_brut / referinta_alba_filtru.",
        "Spectrometrul este citit pe ferestre de +/-10 nm in jurul filtrului, folosind mediana ca sa reduca spike-urile.",
    ]
    for item in bullets:
        bullet(doc, item)

    add_heading(doc, "3. Compensarea curbei sursei", 1)
    add_callout(
        doc,
        "Formula folosita",
        (
            "net_source(lambda) = source(lambda) - dark(lambda)\n"
            "camera_compensata(lambda) = camera_alb(lambda) * net_source(850) / net_source(lambda)\n"
            "greutate_incredere(lambda) = min(1, net_source(lambda) / net_source(850))"
        ),
    )
    table = doc.add_table(rows=1, cols=4)
    headers = ["Valoare", "850 nm", "940 nm", "Interpretare"]
    for cell, text in zip(table.rows[0].cells, headers):
        cell.text = text
    source_940 = data["source_940"]
    corrected_940 = data["corrected_940"]
    rows = [
        (
            "net_source",
            "12292.23",
            f"{source_940['net_source_window_intensity']}",
            "940 are semnal de sursa mult mai slab.",
        ),
        (
            "factor compensare",
            "1.00x",
            f"{source_940['compensation_multiplier_vs_850']}x",
            "940 trebuie amplificat puternic, deci creste si zgomotul.",
        ),
        (
            "greutate incredere",
            "1.00",
            f"{source_940['reliability_weight']}",
            "940 este util ca explorare, nu ca proba principala.",
        ),
        (
            "proba curenta",
            "-",
            f"{corrected_940['white_normalized_camera']} -> {corrected_940['source_curve_compensated_camera']}",
            "Valoarea compensata este salvata, dar marcata cu incredere mica.",
        ),
    ]
    for item in rows:
        cells = table.add_row().cells
        for cell, text in zip(cells, item):
            cell.text = str(text)
    set_widths(table, [4.0, 3.0, 4.1, 5.4])
    style_table(table, header_fill="2D6A4F")

    add_heading(doc, "4. Ce se salveaza acum", 1)
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Fisier / zona"
    table.rows[0].cells[1].text = "Ce contine"
    rows = [
        ("mapping_report.json", "Raport complet: sursa, masca, benzi, indici, avertismente si calibrare."),
        ("camera_calibration.csv", "Valori brute, normalizare cu alb, compensare sursa, greutate de incredere si estimari spectrometru."),
        ("vegetation_mask.png", "Masca frunzei folosita efectiv la mediere."),
        ("average_spectrum.png", "Grafic rapid cu raspunsul mediu pe filtre."),
        ("app text", "Explicatii in romana despre corectie, compensare si avertismente."),
    ]
    for item in rows:
        cells = table.add_row().cells
        for cell, text in zip(cells, item):
            cell.text = text
    set_widths(table, [5.2, 11.3])
    style_table(table)

    add_heading(doc, "5. Reguli de interpretare", 1)
    for item in [
        "NDVI, NDRE si GNDVI sunt mai potrivite pentru concluzii decat NDWI_850_940 in setul actual.",
        "NDWI_850_940 se calculeaza ca valoare raw, dar app-ul il exclude din concluziile calibrate deoarece depinde de 940 nm.",
        "Cand o banda are factor mare de compensare, valoarea poate fi interesanta, dar nu trebuie prezentata ca masuratoare foarte sigura.",
        "Cel mai puternic argument al proiectului este ca semnalul camerei urmareste structura spectrala dupa referinte si calibrare.",
    ]:
        bullet(doc, item)

    add_heading(doc, "6. Verificari facute", 1)
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Verificare"
    table.rows[0].cells[1].text = "Rezultat"
    rows = [
        ("Profil calibrare reconstruit", str(data["profile_id"])),
        ("Rulare pe proba curenta", "Output regenerat in user_check_healthy_yellow_2026_05_24_fixed_mask."),
        ("Teste automate", str(data["pytest"])),
        ("App", "Restartat dupa schimbari ca sa incarce noua calibrare."),
    ]
    for item in rows:
        cells = table.add_row().cells
        for cell, text in zip(cells, item):
            cell.text = text
    set_widths(table, [5.0, 11.5])
    style_table(table, header_fill="6C757D")

    add_heading(doc, "7. Componente atinse", 1)
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Componenta"
    table.rows[0].cells[1].text = "Rol"
    rows = [
        ("core/vegetation.py", "Algoritmul nou de detectie a frunzei."),
        ("analysis_app/import_wizard.py", "Editarea inputului existent si inlocuirea imaginilor individuale."),
        ("core/camera_calibration.py", "Normalizare cu alb, modele spectrometru, compensarea curbei sursei si avertismente."),
        ("core/outputs.py", "Coloane noi in camera_calibration.csv."),
        ("analysis_app/readable.py", "Text citibil in app pentru benzile compensate si greutatea de incredere."),
        ("camera_calibration_profile.json", "Profil reconstruit din datele reale si referintele spectrale."),
    ]
    for item in rows:
        cells = table.add_row().cells
        for cell, text in zip(cells, item):
            cell.text = text
    set_widths(table, [6.0, 10.5])
    style_table(table, header_fill="495057")

    footer = doc.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = footer.add_run("Rezumat intern - app analiza frunze camera + spectrometru")
    r.font.size = Pt(8)
    r.font.color.rgb = RGBColor(110, 110, 110)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build_doc())
