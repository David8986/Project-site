"""Add the real spectrometer/camera validation data to the final Romanian DOCX."""

from __future__ import annotations

import json
import shutil
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
DOCX_PATH = Path(r"C:\Users\david\OneDrive\Desktop\referat modificat FINAL.docx")
BACKUP_PATH = Path(r"C:\Users\david\OneDrive\Desktop\referat modificat FINAL - backup inainte date reale.docx")
SUMMARY_PATH = ROOT / "outputs" / "inside_outside_comparison" / "inside_outside_summary.json"
GRAPHS = ROOT / "outputs" / "inside_outside_comparison" / "graphs"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=100, bottom=90, end=100) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def style_table(table) -> None:
    try:
        table.style = "Table Grid"
    except KeyError:
        pass
    for row_index, row in enumerate(table.rows):
        for cell in row.cells:
            set_cell_margins(cell)
            if row_index == 0:
                set_cell_shading(cell, "DDEAD7")
            else:
                set_cell_shading(cell, "FFFFFF" if row_index % 2 else "F6FAF4")
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.name = "Aptos"
                    run.font.size = Pt(9)
                    if row_index == 0:
                        run.font.bold = True


def set_widths(table, widths_cm: list[float]) -> None:
    for row in table.rows:
        for cell, width in zip(row.cells, widths_cm):
            cell.width = Cm(width)


def add_before(document: Document, ref_paragraph, kind: str, *args, **kwargs):
    """Create a block at document end, then move it before ref_paragraph."""

    if kind == "paragraph":
        block = document.add_paragraph(*args, **kwargs)
        ref_paragraph._p.addprevious(block._p)
        return block
    if kind == "table":
        block = document.add_table(*args, **kwargs)
        ref_paragraph._p.addprevious(block._tbl)
        return block
    raise ValueError(kind)


def format_paragraph(paragraph, *, bold=False, size=11, align=None, before=0, after=6) -> None:
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    if align is not None:
        paragraph.alignment = align
    for run in paragraph.runs:
        run.font.name = "Aptos"
        run.font.size = Pt(size)
        run.font.bold = bold


def add_heading_before(document: Document, ref, text: str) -> None:
    paragraph = add_before(document, ref, "paragraph", text)
    format_paragraph(paragraph, bold=True, size=14, before=10, after=6)


def add_body_before(document: Document, ref, text: str) -> None:
    paragraph = add_before(document, ref, "paragraph", text)
    paragraph.paragraph_format.first_line_indent = Cm(0.8)
    format_paragraph(paragraph, size=11, before=0, after=6)


def add_caption_before(document: Document, ref, text: str) -> None:
    paragraph = add_before(document, ref, "paragraph", text)
    format_paragraph(paragraph, bold=True, size=10.5, align=WD_ALIGN_PARAGRAPH.CENTER, before=8, after=4)


def add_picture_before(document: Document, ref, image_path: Path, width_inches: float = 6.25) -> None:
    paragraph = add_before(document, ref, "paragraph")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    run.add_picture(str(image_path), width=Inches(width_inches))
    paragraph.paragraph_format.space_after = Pt(4)


def update_heading_numbers(document: Document) -> None:
    replacements = {
        "3.4. Concluzia masuratorilor și a proiectului": "3.5. Concluzia măsurătorilor și a proiectului",
        "3.5. Concluzii și Contribuții Originale": "3.6. Concluzii și contribuții originale",
        "3.1. Bibliografie": "4. Bibliografie",
    }
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text in replacements:
            paragraph.text = replacements[text]


def find_insert_reference(document: Document):
    for paragraph in document.paragraphs:
        if paragraph.text.strip().startswith("3.4. Concluzia"):
            return paragraph
    raise RuntimeError("Nu am gasit heading-ul 3.4 pentru inserare.")


def add_validation_section(document: Document) -> None:
    summary = json.loads(SUMMARY_PATH.read_text(encoding="utf-8"))
    ref = find_insert_reference(document)

    add_heading_before(document, ref, "3.4. Validarea experimentală prin comparație spectrometrică")
    add_body_before(
        document,
        ref,
        (
            "Pentru a verifica practic sistemul, măsurătorile spectrometrice realizate afară au fost "
            "comparate cu măsurătorile realizate în cutia de testare. Comparația a fost făcută pe perechi "
            "cu aceeași stare a frunzei și același filtru spectral, folosind întregul spectru măsurat."
        ),
    )
    add_body_before(
        document,
        ref,
        (
            "Scopul acestei etape nu este ca valorile brute să fie identice, deoarece iluminarea și geometria "
            "măsurării sunt diferite. Important este dacă forma curbei și tendințele dintre filtre se păstrează "
            "suficient de bine pentru ca sistemul optic cu cameră să poată fi calibrat."
        ),
    )

    table = add_before(document, ref, "table", rows=1, cols=3)
    headers = ["Indicator", "Valoare", "Interpretare"]
    for cell, text in zip(table.rows[0].cells, headers):
        cell.text = text
    rows = [
        (
            "Perechi spectrale comparate",
            str(summary["matched_inside_outside_spectrum_pairs"]),
            "măsurători afară corelate cu măsurători în cutie",
        ),
        (
            "Similaritate mediană totală",
            f"r = {summary['overall_median_shape_similarity_r']:.2f}".replace(".", ","),
            "asemănare moderată a formei spectrale",
        ),
        (
            "Similaritate frunze sănătoase",
            f"r = {summary['by_condition']['healthy']['median_shape_similarity_r']:.2f}".replace(".", ","),
            "rezultat apropiat de media generală",
        ),
        (
            "Similaritate frunze nesănătoase",
            f"r = {summary['by_condition']['unhealthy']['median_shape_similarity_r']:.2f}".replace(".", ","),
            "asemănare moderată, cu variații mai mari",
        ),
        (
            "Raport median al semnalului",
            f"{summary['overall_median_outside_to_inside_area_ratio']:.1f}x".replace(".", ","),
            "semnalul brut este mult mai puternic afară",
        ),
    ]
    for item in rows:
        cells = table.add_row().cells
        for cell, text in zip(cells, item):
            cell.text = text
    set_widths(table, [5.0, 3.0, 8.2])
    style_table(table)

    figures = [
        (
            "Figura 1. Compararea formei spectrelor normalizate",
            GRAPHS / "01_mean_normalized_spectra_inside_vs_outside.png",
            (
                "Spectrele sunt normalizate pentru ca diferența de intensitate a luminii să nu domine comparația. "
                "Se observă că mediul de măsurare schimbă forma exactă a curbei, dar păstrează o tendință generală "
                "utilă pentru calibrarea camerei."
            ),
        ),
        (
            "Figura 2. Similaritatea pe filtre și pe starea frunzei",
            GRAPHS / "03_similarity_by_filter_and_condition.png",
            (
                "Coeficientul Pearson r arată cât de apropiată este forma spectrului măsurat în cutie față de "
                "cel măsurat afară. Cele mai utile rezultate apar în zona red-edge și NIR, unde sistemul păstrează "
                "o asemănare moderată sau bună."
            ),
        ),
        (
            "Figura 3. Diferența dintre intensitatea brută afară și în cutie",
            GRAPHS / "04_raw_signal_area_inside_vs_outside.png",
            (
                "Aria totală a semnalului spectral este mult mai mare afară decât în cutie. Raportul median de "
                "aproximativ 137,8x arată că valorile brute nu trebuie comparate direct, ci după normalizare."
            ),
        ),
        (
            "Figura 4. Zgomotul și rugozitatea curbelor spectrale",
            GRAPHS / "05_noise_roughness_inside_vs_outside.png",
            (
                "În cutie, semnalul este mai slab, iar variațiile mici devin mai vizibile. Acest lucru nu invalidează "
                "metoda, dar arată că sistemul are nevoie de expunere controlată, calibrare și repetarea măsurătorilor."
            ),
        ),
        (
            "Figura 5. Harta de similaritate pentru perechile spectrale",
            GRAPHS / "02_inside_outside_similarity_heatmap.png",
            (
                "Culorile verzi indică perechi cu asemănare mai bună, iar zonele galben-roșii indică perechi mai "
                "sensibile la condițiile de măsurare. Mediana r = 0,56 arată o legătură măsurabilă, utilă pentru "
                "calibrarea aplicației."
            ),
        ),
    ]
    for caption, image, explanation in figures:
        add_caption_before(document, ref, caption)
        add_picture_before(document, ref, image)
        add_body_before(document, ref, explanation)

    add_heading_before(document, ref, "Interpretarea rezultatelor experimentale")
    add_body_before(
        document,
        ref,
        (
            "Rezultatele arată că măsurătorile din cutie și cele din exterior nu sunt identice, însă păstrează o "
            "legătură măsurabilă. Similaritatea mediană de aproximativ r = 0,56 indică o asemănare moderată a "
            "formei spectrale, ceea ce este important pentru un prototip low-cost aflat în etapa de validare."
        ),
    )
    add_body_before(
        document,
        ref,
        (
            "Prin urmare, sistemul nu trebuie evaluat doar prin valori brute de luminozitate, ci prin comparații "
            "normalizate, rapoarte între benzi și comportamentul relativ al curbelor. Această concluzie susține "
            "obiectivul proiectului: realizarea unei camere multispectrale accesibile care poate aproxima informații "
            "utile despre starea plantei și poate fi îmbunătățită prin calibrare."
        ),
    )


def main() -> None:
    if not BACKUP_PATH.exists():
        shutil.copy2(DOCX_PATH, BACKUP_PATH)
    document = Document(DOCX_PATH)
    add_validation_section(document)
    update_heading_numbers(document)
    document.save(DOCX_PATH)
    print(DOCX_PATH)
    print(BACKUP_PATH)


if __name__ == "__main__":
    main()
