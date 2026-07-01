"""Build a consolidated Excel workbook for the leaf photo/spectrometer analysis."""

from __future__ import annotations

import csv
import math
import re
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


PROJECT_ROOT = Path(r"C:\Users\david\OneDrive\Desktop\Archive\Projects\CodeX")
DATASET = Path(r"C:\Users\david\OneDrive\Desktop\Archive\Misc\Review Folders\New folder (6)")
ANALYSIS_DIR = PROJECT_ROOT / "outputs" / "spectrometer_photo_comparison"
OUTPUT_XLSX = ANALYSIS_DIR / "leaf_photo_spectrometer_numeric_analysis.xlsx"


def main() -> None:
    ensure_source_outputs()

    wb = Workbook()
    default = wb.active
    wb.remove(default)

    overview_rows = build_overview_rows()
    add_rows_sheet(wb, "Overview", overview_rows)
    add_csv_sheet(wb, "Photo Raw Metrics", ANALYSIS_DIR / "per_photo_leaf_analysis.csv")
    add_csv_sheet(wb, "Photo Group Summary", ANALYSIS_DIR / "photo_group_summary.csv")
    add_csv_sheet(wb, "Photo Associations", ANALYSIS_DIR / "photo_to_filter_association.csv")
    add_csv_sheet(wb, "Spectrometer Catalog", ANALYSIS_DIR / "spectrometer_file_catalog.csv")
    add_rows_sheet(wb, "OCV Numeric Settings", extract_ocv_numeric_settings())
    add_csv_sheet(wb, "All Pair Associations", ANALYSIS_DIR / "proposed_measurement_associations.csv")

    add_summary_chart(wb)
    format_workbook(wb)
    ANALYSIS_DIR.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT_XLSX)
    print(OUTPUT_XLSX)


def ensure_source_outputs() -> None:
    # The analysis scripts should already have been run in the prior step. Keep
    # this explicit so the workbook fails loudly if a required source is missing.
    required = [
        ANALYSIS_DIR / "per_photo_leaf_analysis.csv",
        ANALYSIS_DIR / "photo_group_summary.csv",
        ANALYSIS_DIR / "photo_to_filter_association.csv",
        ANALYSIS_DIR / "spectrometer_file_catalog.csv",
        ANALYSIS_DIR / "proposed_measurement_associations.csv",
    ]
    missing = [path for path in required if not path.exists()]
    if missing:
        raise FileNotFoundError("Missing analysis CSV files: " + ", ".join(str(path) for path in missing))


def build_overview_rows() -> list[dict[str, object]]:
    photo_rows = read_csv_dicts(ANALYSIS_DIR / "per_photo_leaf_analysis.csv")
    summary_rows = read_csv_dicts(ANALYSIS_DIR / "photo_group_summary.csv")
    catalog_rows = read_csv_dicts(ANALYSIS_DIR / "spectrometer_file_catalog.csv")
    ocv_settings = extract_ocv_numeric_settings()
    non_duplicate_ocv = [
        row for row in catalog_rows if "New folder" not in str(row.get("relative_path", "")).split("\\")
    ]

    rows: list[dict[str, object]] = [
        {"metric": "dataset_path", "value": str(DATASET), "notes": "Source folder analyzed"},
        {"metric": "photo_count", "value": len(photo_rows), "notes": "JPG files analyzed"},
        {"metric": "photo_group_count", "value": len(summary_rows), "notes": "Camera Roll photo folders"},
        {"metric": "ocv_file_count_cataloged", "value": len(catalog_rows), "notes": "Includes duplicate top-level/New folder copies"},
        {"metric": "ocv_file_count_non_duplicate", "value": len(non_duplicate_ocv), "notes": "Excludes paths under duplicate New folder"},
        {
            "metric": "ocv_numeric_setting_rows",
            "value": len(ocv_settings),
            "notes": "Numeric settings extracted from OceanView XML entries",
        },
        {
            "metric": "raw_spectral_intensity_arrays_found",
            "value": 0,
            "notes": "No CSV/TXT/wavelength-intensity arrays found inside checked OCV containers",
        },
    ]
    for summary in summary_rows:
        group = summary.get("photo_group", "")
        for key in ("avg_gcc", "avg_exg", "avg_vari", "avg_mean_saturation", "avg_mean_value_brightness"):
            rows.append(
                {
                    "metric": f"{group}_{key}",
                    "value": to_number(summary.get(key, "")),
                    "notes": "Computed from segmented leaf pixels in JPG photos",
                }
            )
    return rows


def extract_ocv_numeric_settings() -> list[dict[str, object]]:
    rows: list[dict[str, object]] = []
    for path in sorted(DATASET.rglob("*.ocv")):
        if "\\New folder\\" in str(path):
            continue
        if not zipfile.is_zipfile(path):
            continue
        metadata = infer_measurement_metadata(path)
        with zipfile.ZipFile(path) as zf:
            for entry in zf.infolist():
                if not entry.filename.lower().endswith(".xml"):
                    continue
                text = zf.read(entry).decode("utf-8", errors="ignore")
                rows.extend(extract_numeric_xml_rows(path, entry.filename, text, metadata))
    return rows


def extract_numeric_xml_rows(path: Path, entry_name: str, text: str, metadata: dict[str, str]) -> list[dict[str, object]]:
    rows: list[dict[str, object]] = []
    try:
        root = ET.fromstring(text)
    except ET.ParseError:
        return rows

    key = text_of(root.find("key")) or root.attrib.get("name", "")
    value = text_of(root.find("value"))
    numeric = to_number(value)
    if isinstance(numeric, (int, float)) and math.isfinite(float(numeric)):
        rows.append(
            {
                "ocv_file": path.name,
                "relative_path": str(path.relative_to(DATASET)),
                "zip_entry": entry_name,
                "setting_key": key,
                "numeric_value": numeric,
                "leaf_condition": metadata["leaf_condition"],
                "measurement_type": metadata["measurement_type"],
                "filter_nm": metadata["filter_nm"],
            }
        )

    for tag in (
        "activeXMinAxis",
        "activeXMaxAxis",
        "activeYMinAxis",
        "activeYMaxAxis",
        "cursorXPoint",
        "xMinAxis",
        "xMaxAxis",
        "yMinAxis",
        "yMaxAxis",
        "runDelay",
        "pointX",
        "pointY",
    ):
        element = root.find(tag)
        number = to_number(text_of(element))
        if isinstance(number, (int, float)) and math.isfinite(float(number)):
            rows.append(
                {
                    "ocv_file": path.name,
                    "relative_path": str(path.relative_to(DATASET)),
                    "zip_entry": entry_name,
                    "setting_key": tag,
                    "numeric_value": number,
                    "leaf_condition": metadata["leaf_condition"],
                    "measurement_type": metadata["measurement_type"],
                    "filter_nm": metadata["filter_nm"],
                }
            )
    return rows


def infer_measurement_metadata(path: Path) -> dict[str, str]:
    stem_text = path.stem.lower()
    parent_text = path.parent.name.lower()
    text = f"{parent_text} {stem_text}"
    wavelength = ""
    for value in (532, 556, 680, 725, 850, 940):
        if re.search(rf"(^|[^0-9]){value}([^0-9]|$)", text):
            wavelength = str(value)
            break

    if "maro" in stem_text:
        condition = "brown leaf area"
    elif "galben" in stem_text or "_g_" in stem_text or stem_text.startswith("f_g"):
        condition = "yellow/unhealthy leaf"
    elif "verde" in stem_text or "_v_" in stem_text or stem_text.startswith("f_v"):
        condition = "green/healthy leaf"
    elif any(token in text for token in ("galben", "_g_", "f_g", "nesanatoasa")):
        condition = "yellow/unhealthy leaf"
    elif any(token in text for token in ("verde", "_v_", "f_v")):
        condition = "green/healthy leaf"
    elif "uvvis" in text or "uv vis" in text:
        condition = "outside/UV-VIS reference"
    else:
        condition = "unknown"

    if "spectru" in text and "intreg" in text:
        measurement = "full spectrum"
    elif "fara filtru" in text or "farafiltru" in text or "direct" in text:
        measurement = "direct/no filter"
    elif wavelength:
        measurement = "filtered"
    else:
        measurement = "unknown"

    return {"leaf_condition": condition, "measurement_type": measurement, "filter_nm": wavelength}


def add_csv_sheet(wb: Workbook, name: str, path: Path) -> None:
    rows = read_csv_dicts(path)
    add_rows_sheet(wb, name, rows)


def add_rows_sheet(wb: Workbook, name: str, rows: list[dict[str, object]]) -> None:
    ws = wb.create_sheet(name)
    if not rows:
        ws.append(["No data"])
        return
    headers = list(rows[0].keys())
    ws.append(headers)
    for row in rows:
        ws.append([to_number(row.get(header, "")) for header in headers])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def add_summary_chart(wb: Workbook) -> None:
    ws = wb["Photo Group Summary"]
    chart = BarChart()
    chart.type = "col"
    chart.style = 10
    chart.title = "Photo Group Numeric Metrics"
    chart.y_axis.title = "Metric value"
    chart.x_axis.title = "Photo group"
    data = Reference(ws, min_col=14, max_col=20, min_row=1, max_row=3)
    cats = Reference(ws, min_col=1, min_row=2, max_row=3)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    chart.height = 8
    chart.width = 18
    ws.add_chart(chart, "W2")


def format_workbook(wb: Workbook) -> None:
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    for ws in wb.worksheets:
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=False)
        for column_cells in ws.columns:
            column_letter = get_column_letter(column_cells[0].column)
            max_len = max(len(str(cell.value)) if cell.value is not None else 0 for cell in column_cells[:200])
            ws.column_dimensions[column_letter].width = min(max(max_len + 2, 10), 48)
        ws.row_dimensions[1].height = 30


def read_csv_dicts(path: Path) -> list[dict[str, object]]:
    with path.open(newline="", encoding="utf-8") as handle:
        return list(csv.DictReader(handle))


def text_of(element: object) -> str:
    if element is None:
        return ""
    text = getattr(element, "text", None)
    return text.strip() if text else ""


def to_number(value: object) -> object:
    if value is None:
        return ""
    if isinstance(value, (int, float)):
        return value
    text = str(value).strip()
    if text == "":
        return ""
    try:
        number = float(text)
    except ValueError:
        return value
    if math.isfinite(number) and number.is_integer():
        return int(number)
    return number


if __name__ == "__main__":
    main()
