"""Build an Excel workbook from the real leaf analysis CSV outputs."""

from __future__ import annotations

import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


OUT_DIR = Path(r"C:\Users\david\OneDrive\Desktop\Archive\Projects\CodeX\outputs\real_leaf_analysis")
WORKBOOK_PATH = OUT_DIR / "real_leaf_camera_spectrometer_analysis.xlsx"


def main() -> None:
    wb = Workbook()
    wb.remove(wb.active)
    add_rows_sheet(wb, "Summary", build_summary_rows())
    add_csv_sheet(wb, "Spectra Features", OUT_DIR / "spectra_features.csv")
    add_csv_sheet(wb, "Camera Features", OUT_DIR / "camera_features.csv")
    add_csv_sheet(wb, "Camera Spectrum Pairs", OUT_DIR / "camera_spectra_pairs.csv")
    add_csv_sheet(wb, "Spectra Long Sampled", OUT_DIR / "spectra_long_sampled.csv", max_rows=5000)
    add_charts(wb)
    format_workbook(wb)
    wb.save(WORKBOOK_PATH)
    print(WORKBOOK_PATH)


def build_summary_rows() -> list[dict[str, object]]:
    spectra = read_csv(OUT_DIR / "spectra_features.csv")
    cameras = read_csv(OUT_DIR / "camera_features.csv")
    pairs = read_csv(OUT_DIR / "camera_spectra_pairs.csv")
    samples = [row for row in spectra if row["role"] == "sample"]
    rows: list[dict[str, object]] = [
        {"metric": "spectra_files_parsed", "value": len(spectra), "notes": "TXT spectra parsed"},
        {"metric": "sample_spectra", "value": len(samples), "notes": "Leaf/sample spectra only"},
        {"metric": "camera_images", "value": len(cameras), "notes": "JPG images analyzed"},
        {"metric": "paired_camera_spectrum_rows", "value": len(pairs), "notes": "Same sample and filter/wavelength"},
        {"metric": "reflectance_ok_count", "value": sum(1 for row in samples if row["reflectance_status"] == "ok"), "notes": "Reference calibration completed"},
    ]
    for condition in sorted({row["condition"] for row in samples}):
        subset = [row for row in samples if row["condition"] == condition]
        rows.append({"metric": f"{condition}_sample_spectra", "value": len(subset), "notes": ""})
    return rows


def add_csv_sheet(wb: Workbook, name: str, path: Path, max_rows: int | None = None) -> None:
    rows = read_csv(path)
    if max_rows is not None:
        rows = rows[:max_rows]
    add_rows_sheet(wb, name, rows)


def add_rows_sheet(wb: Workbook, name: str, rows: list[dict[str, object]]) -> None:
    ws = wb.create_sheet(name)
    if not rows:
        ws.append(["No data"])
        return
    headers = list(rows[0].keys())
    ws.append(headers)
    for row in rows:
        ws.append([coerce(row.get(header, "")) for header in headers])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def add_charts(wb: Workbook) -> None:
    ws = wb["Summary"]
    chart = BarChart()
    chart.title = "Analysis Row Counts"
    chart.y_axis.title = "Count"
    data = Reference(ws, min_col=2, min_row=1, max_row=5)
    cats = Reference(ws, min_col=1, min_row=2, max_row=5)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    chart.height = 7
    chart.width = 13
    ws.add_chart(chart, "E2")

    pairs = wb["Camera Spectrum Pairs"]
    if pairs.max_row > 2:
        scatter_like = LineChart()
        scatter_like.title = "Camera Brightness by Pair Row"
        scatter_like.y_axis.title = "Brightness"
        scatter_like.x_axis.title = "Pair row"
        data = Reference(pairs, min_col=8, min_row=1, max_row=min(pairs.max_row, 30))
        scatter_like.add_data(data, titles_from_data=True)
        scatter_like.height = 7
        scatter_like.width = 13
        pairs.add_chart(scatter_like, "R2")


def format_workbook(wb: Workbook) -> None:
    fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    for ws in wb.worksheets:
        for cell in ws[1]:
            cell.fill = fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for col in ws.columns:
            letter = get_column_letter(col[0].column)
            max_len = max(len(str(cell.value)) if cell.value is not None else 0 for cell in col[:200])
            ws.column_dimensions[letter].width = min(max(max_len + 2, 10), 42)
        ws.row_dimensions[1].height = 32


def read_csv(path: Path) -> list[dict[str, object]]:
    with path.open(newline="", encoding="utf-8") as handle:
        return list(csv.DictReader(handle))


def coerce(value: object) -> object:
    if value is None:
        return ""
    text = str(value)
    if text == "":
        return ""
    try:
        number = float(text)
    except ValueError:
        return value
    if number.is_integer():
        return int(number)
    return number


if __name__ == "__main__":
    main()
