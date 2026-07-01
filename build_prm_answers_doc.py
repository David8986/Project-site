from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUT = r"D:\downloads\Rezolvare_PRM_exercitii_comune_fara_eseu.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(text)
    run.bold = bold
    run.font.name = "Aptos"
    run.font.size = Pt(10)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_answer(doc, number, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(4)
    p.add_run(text)


def add_section_note(doc, text):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_shading(cell, "EAF2F8")
    set_cell_text(cell, text)
    for paragraph in cell.paragraphs:
        paragraph.paragraph_format.left_indent = Cm(0.1)
        paragraph.paragraph_format.right_indent = Cm(0.1)
        paragraph.paragraph_format.space_before = Pt(4)
        paragraph.paragraph_format.space_after = Pt(4)
    doc.add_paragraph()


def add_pair_table(doc, headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        set_cell_shading(hdr[i], "D9EAF7")
        set_cell_text(hdr[i], h, bold=True)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], value)
    doc.add_paragraph()


def style_document(doc):
    section = doc.sections[0]
    section.top_margin = Cm(1.7)
    section.bottom_margin = Cm(1.7)
    section.left_margin = Cm(1.9)
    section.right_margin = Cm(1.9)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.08
    normal.paragraph_format.space_after = Pt(6)

    for name, size, color in [
        ("Title", 18, "1F4E79"),
        ("Heading 1", 14, "1F4E79"),
        ("Heading 2", 12, "365F91"),
    ]:
        style = styles[name]
        style.font.name = "Aptos Display" if name == "Title" else "Aptos"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(10 if name != "Title" else 0)
        style.paragraph_format.space_after = Pt(6)


doc = Document()
style_document(doc)

title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.add_run("Rezolvare - Primul Razboi Mondial")

subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle.add_run("Exercitii comune identificate si raspunsuri, fara eseu").italic = True

doc.add_heading("Exercitii comune", level=1)
add_section_note(
    doc,
    "Nu trebuie rezolvate de doua ori: textul cu francezii care credeau ca ajung la Berlin si textul despre transee apar in DOCX Varianta 3 si in poza 1 Varianta 2."
)
for item in [
    "Revolutia bolsevica din 1917 apare in mai multe variante: raspunsul este Rusia.",
    "Intrebarea despre regele Romaniei in Primul Razboi Mondial: raspunsul este Ferdinand I.",
    "Intrebarea despre Tratatele de la Paris-Versailles se bazeaza pe recunoasterea Marii Uniri si pe formarea Romaniei Mari.",
]:
    p = doc.add_paragraph(style="List Bullet")
    p.add_run(item)

doc.add_heading("DOCX - Varianta 3", level=1)
doc.add_heading("I. Grila", level=2)
add_pair_table(
    doc,
    ["Item", "Raspuns"],
    [
        ["1", "d) Balcanica"],
        ["2", "c) Japonia"],
        ["3", "b) turci"],
        ["4", "c) Ferdinand I"],
        ["5", "b) Rusia"],
    ],
)

doc.add_heading("II. Corespondenta", level=2)
add_pair_table(
    doc,
    ["Item", "Raspuns"],
    [["1", "c"], ["2", "d"], ["3", "e"], ["4", "a"], ["5", "b"]],
)

doc.add_heading("III. Argumente", level=2)
for text in [
    "In 1914 Romania s-a declarat neutra deoarece nu era pregatita sa intre imediat in razboi, iar clasa politica era impartita intre sustinatorii Antantei si cei ai Puterilor Centrale. Neutralitatea ii permitea Romaniei sa astepte momentul favorabil pentru a-si urmari obiectivele nationale.",
    "In august 1916 Romania a intrat in razboi de partea Antantei deoarece Antanta promitea sprijin pentru unirea teritoriilor romanesti aflate sub stapanire austro-ungara, mai ales Transilvania. De aceea armata romana a atacat Austro-Ungaria prin trecerea Carpatilor.",
]:
    p = doc.add_paragraph(style="List Number")
    p.add_run(text)

doc.add_heading("IV. Textul A/B", level=2)
for text in [
    "Francezii s-au gandit la un razboi scurt, rapid, de miscare.",
    "O batalie de pe frontul de vest a fost Verdun, desfasurata in 1916 intre francezi si germani. A fost una dintre cele mai grele batalii ale razboiului, cu pierderi foarte mari, iar francezii au reusit sa reziste ofensivei germane.",
    "Doua informatii din textul B: soldatii stateau mult timp in transee fara sa se poata spala; conditiile erau groaznice, cu noroi, apa din gropi de obuz si boli precum «picioare de transee» sau «febra de transee».",
    "O cauza a infrangerii Germaniei a fost intrarea SUA in razboi de partea Antantei, ceea ce a oferit aliatilor resurse militare si economice superioare.",
    "O consecinta a Tratatelor de la Paris-Versailles a fost redesenarea hartii Europei si recunoasterea unor noi state sau granite. Pentru Romania, tratatele au contribuit la recunoasterea internationala a Marii Uniri.",
]:
    p = doc.add_paragraph(style="List Number")
    p.add_run(text)

doc.add_heading("Poza 1 - Varianta 2", level=1)
doc.add_heading("I. Grila", level=2)
add_pair_table(
    doc,
    ["Item", "Raspuns"],
    [
        ["1", "c) Serbia"],
        ["2", "b) Bulgaria"],
        ["3", "d) 1916"],
        ["4", "c) Ferdinand I"],
        ["5", "c) Rusia"],
    ],
)

doc.add_heading("II. Sursa despre intrarea Romaniei in razboi", level=2)
for text in [
    "Regele Romaniei era Ferdinand I.",
    "Data intrarii Romaniei in razboi: 15/28 august 1916, conform sursei.",
    "O cauza: dorinta de a-i elibera pe romanii ardeleni aflati sub ocupatie austro-ungara.",
    "O actiune militara: armata romana a trecut Carpatii si a atacat Austro-Ungaria in Transilvania, dupa intrarea Romaniei in razboi de partea Antantei.",
    "Tratatele de la Paris-Versailles au fost importante pentru Romania deoarece au recunoscut pe plan international Marea Unire din 1918 si noile granite ale statului roman.",
]:
    p = doc.add_paragraph(style="List Number")
    p.add_run(text)

doc.add_heading("III. Textul A/B", level=2)
add_section_note(doc, "Acesta este exercitiul comun cu DOCX Varianta 3, sectiunea IV. Foloseste raspunsurile de acolo.")
for text in [
    "Hotarare: recunoasterea unirii Transilvaniei, Bucovinei si Basarabiei cu Romania prin tratatele de pace.",
    "Consecinta: formarea Romaniei Mari si consolidarea statului national unitar roman.",
]:
    p = doc.add_paragraph(style="List Bullet")
    p.add_run(text)

doc.add_heading("Poza 2 - Fara eseu", level=1)
doc.add_heading("I. Sursa despre diplomatia Romaniei interbelice", level=2)
for text in [
    "Conflictul militar international mentionat este Primul Razboi Mondial.",
    "Secolul la care se refera sursa este secolul al XX-lea.",
    "Diplomatul roman este Nicolae Titulescu; el a fost presedinte al Adunarii Generale a Societatii Natiunilor.",
    "Doua informatii despre relatiile cu Uniunea Sovietica: Romania a reluat relatiile diplomatice cu URSS in 1934; Titulescu a purtat convorbiri cu Litvinov pentru un tratat de asistenta mutuala.",
    "Punct de vedere: In perioada 1921-1934, Romania a dus o politica externa de aparare a granitelor si de mentinere a sistemului de la Versailles. Acest lucru este sustinut de participarea la Mica Intelegere in 1921 si de constituirea Intelegerii Balcanice in 1934.",
    "Conferinta de la Paris-Versailles a fost importanta pentru Romania deoarece a dus la recunoasterea internationala a Marii Uniri. De exemplu, prin tratatele de pace au fost confirmate unirea Transilvaniei, Bucovinei si Basarabiei cu Romania, ceea ce a consolidat Romania Mare.",
]:
    p = doc.add_paragraph(style="List Number")
    p.add_run(text)

doc.core_properties.title = "Rezolvare PRM - exercitii comune si raspunsuri"
doc.core_properties.subject = "Istorie clasa a X-a"
doc.save(OUT)
print(OUT)
