from docx import Document
from docx.shared import Cm, Pt


OUT = r"D:\downloads\Rezolvare_PRM_scurt_copy_paste.docx"


def setup(doc):
    section = doc.sections[0]
    section.top_margin = Cm(1.6)
    section.bottom_margin = Cm(1.6)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)
    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(3)
    for name, size in [("Title", 16), ("Heading 1", 13), ("Heading 2", 11)]:
        style = doc.styles[name]
        style.font.name = "Aptos"
        style.font.size = Pt(size)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(8)
        style.paragraph_format.space_after = Pt(4)


def h1(doc, text):
    doc.add_heading(text, level=1)


def h2(doc, text):
    doc.add_heading(text, level=2)


def line(doc, text):
    doc.add_paragraph(text)


doc = Document()
setup(doc)

doc.add_heading("Rezolvare PRM - scurt, copy-paste", level=0)

h1(doc, "Exercitii comune")
line(doc, "Textul A/B despre francezi si transee apare in DOCX Varianta 3 si in poza 1 Varianta 2. Se rezolva o singura data.")
line(doc, "Raspuns comun: Revolutia bolsevica - Rusia; regele Romaniei - Ferdinand I; Paris-Versailles - recunoasterea Marii Uniri.")

h1(doc, "DOCX - Varianta 3")
h2(doc, "I. Grila")
line(doc, "1-d) Balcanica; 2-c) Japonia; 3-b) turci; 4-c) Ferdinand I; 5-b) Rusia.")

h2(doc, "II. Corespondenta")
line(doc, "1-c; 2-d; 3-e; 4-a; 5-b.")

h2(doc, "III. Argumente")
line(doc, "1. Romania s-a declarat neutra in 1914 deoarece nu era pregatita si clasa politica era impartita intre Antanta si Puterile Centrale.")
line(doc, "2. Romania a intrat in razboi in 1916 de partea Antantei pentru unirea teritoriilor romanesti aflate sub stapanire austro-ungara, mai ales Transilvania.")

h2(doc, "IV. Textul A/B")
line(doc, "1. Francezii s-au gandit la un razboi scurt, rapid, de miscare.")
line(doc, "2. Batalia de la Verdun, din 1916, a fost purtata intre francezi si germani si a produs pierderi foarte mari.")
line(doc, "3. Soldatii stateau in transee fara sa se spele; conditiile erau groaznice, cu noroi, apa din gropi de obuz si boli.")
line(doc, "4. O cauza a infrangerii Germaniei a fost intrarea SUA in razboi de partea Antantei.")
line(doc, "5. O consecinta a Tratatelor de la Paris-Versailles a fost recunoasterea unor noi granite; pentru Romania, recunoasterea Marii Uniri.")

h1(doc, "Poza 1 - Varianta 2")
h2(doc, "I. Grila")
line(doc, "1-c) Serbia; 2-b) Bulgaria; 3-d) 1916; 4-c) Ferdinand I; 5-c) Rusia.")

h2(doc, "II. Sursa despre intrarea Romaniei in razboi")
line(doc, "1. Regele Romaniei era Ferdinand I.")
line(doc, "2. Data intrarii Romaniei in razboi: 15/28 august 1916.")
line(doc, "3. Cauza: eliberarea romanilor ardeleni aflati sub ocupatie austro-ungara.")
line(doc, "4. Actiune militara: armata romana a trecut Carpatii si a atacat Austro-Ungaria in Transilvania.")
line(doc, "5. Tratatele de la Paris-Versailles au fost importante deoarece au recunoscut Marea Unire si noile granite ale Romaniei.")

h2(doc, "III. Textul A/B")
line(doc, "Este comun cu DOCX Varianta 3, exercitiul IV.")
line(doc, "Hotarare: recunoasterea unirii Transilvaniei, Bucovinei si Basarabiei cu Romania.")
line(doc, "Consecinta: formarea Romaniei Mari.")

h1(doc, "Poza 2 - fara eseu")
h2(doc, "I. Sursa despre diplomatia Romaniei interbelice")
line(doc, "1. Conflictul mentionat este Primul Razboi Mondial.")
line(doc, "2. Secolul este al XX-lea.")
line(doc, "3. Diplomatul roman este Nicolae Titulescu; functie: presedinte al Adunarii Generale a Societatii Natiunilor.")
line(doc, "4. Romania a reluat relatiile diplomatice cu URSS in 1934; Titulescu a purtat convorbiri cu Litvinov pentru un tratat de asistenta mutuala.")
line(doc, "5. Romania a urmarit apararea granitelor si mentinerea sistemului de la Versailles, prin Mica Intelegere si Intelegerea Balcanica.")
line(doc, "6. Conferinta de la Paris-Versailles a recunoscut Marea Unire, confirmand unirea Transilvaniei, Bucovinei si Basarabiei cu Romania.")

doc.save(OUT)
print(OUT)
